"""PPTX extraction with python-pptx.

Each slide is one ``Page`` numbered from 1. Shapes are read top to bottom, then left to right,
descending into groups. A shape's frame (converted from EMU to points) is the ``bbox`` of every
span taken from it, so the paragraphs of one text box share a box. Shapes inside a group are
positioned in the group's own coordinate space.
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from pathlib import Path

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from classdeck.ingest.base import (
    ImageSink,
    PageBuilder,
    is_list_text,
    normalize_text,
    table_text,
)
from classdeck.lesson import Page, SpanKind

_EMU_PER_POINT = 12700


class PptxExtractor:
    media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def extract(self, path: Path, *, doc_id: str, image_dir: Path | None) -> list[Page]:
        sink = ImageSink(image_dir, doc_id)
        presentation = Presentation(str(path))
        pages = []
        for number, slide in enumerate(presentation.slides, start=1):
            builder = PageBuilder(number)
            title = slide.shapes.title
            title_id = title.shape_id if title is not None else None
            for shape in _walk(slide.shapes):
                _add_shape(builder, shape, sink, is_title=shape.shape_id == title_id)
            pages.append(builder.build())
        return pages


def _walk(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """Shapes in reading order, with the members of groups in place of the group."""
    for shape in sorted(shapes, key=_position):
        if isinstance(shape, GroupShape):
            yield from _walk(shape.shapes)
        else:
            yield shape


def _position(shape: BaseShape) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def _frame(shape: BaseShape) -> tuple[float, float, float, float] | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    left = shape.left / _EMU_PER_POINT
    top = shape.top / _EMU_PER_POINT
    return (left, top, left + shape.width / _EMU_PER_POINT, top + shape.height / _EMU_PER_POINT)


def _add_shape(builder: PageBuilder, shape: BaseShape, sink: ImageSink, *, is_title: bool) -> None:
    bbox = _frame(shape)
    if isinstance(shape, Picture):
        saved = sink.save(shape.image.blob, page=builder.number, index=builder.image_count + 1)
        builder.add(SpanKind.IMAGE, "", bbox=bbox, image_path=saved)
    elif isinstance(shape, GraphicFrame) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        builder.add(SpanKind.TABLE, table_text(rows), bbox=bbox)
    elif isinstance(shape, Shape) and shape.has_text_frame:
        paragraphs = shape.text_frame.paragraphs
        if is_title:
            title = " ".join(paragraph.text for paragraph in paragraphs)
            builder.add(SpanKind.HEADING, normalize_text(title), bbox=bbox)
            return
        for paragraph in paragraphs:
            text = normalize_text(paragraph.text)
            is_item = paragraph.level > 0 or is_list_text(text)
            builder.add(SpanKind.LIST_ITEM if is_item else SpanKind.PARAGRAPH, text, bbox=bbox)
