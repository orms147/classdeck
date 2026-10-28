"""Word writer: the lesson plan as a printable document, with a source map at the end.

The layout follows the export conventions in ``docs/english-rules.md``: a header, objectives,
teaching aids, then the procedure as a two-column table (Teacher's activities / Students'
activities) grouped by phase with minutes. Competences developed by an activity are highlighted.
Each procedure row ends with a superscript marker ``[n]`` that points at row ``n`` of the source
map in the appendix; blocks the model added on its own also carry an AI-added badge.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.styles.style import ParagraphStyle
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from classdeck.curriculum import Curriculum
from classdeck.export.source_map import (
    AI_ADDED_BADGE,
    PROVENANCE_LABEL,
    build_source_map,
    collect_warnings,
    ensure_exportable,
    format_check,
    origin_text,
)
from classdeck.lesson import (
    EXERCISE_KINDS,
    ActivityBlock,
    Block,
    BlockKind,
    DialogueBlock,
    Lesson,
    ObjectiveBlock,
    Phase,
    PhonicsBlock,
    Provenance,
    Section,
    SentencePatternBlock,
    SourceDocument,
    SummaryBlock,
    TeacherNoteBlock,
    VocabItem,
    VocabularyBlock,
)
from classdeck.lesson.exercises import (
    answer_key,
    exercise_lines,
    exercise_title,
    instruction_of,
)
from classdeck.lesson.labels import lesson_label, unit_line, unit_number, unit_title

FONT = "Arial"
FONT_SIZE = 12
HEADER_FILL = "D9E2F3"
WRAP_UP_FALLBACK = "Review the new words and the pattern at home."

PROCEDURE_WIDTHS = (8.5, 8.5)
APPENDIX_WIDTHS = (0.9, 1.7, 2.6, 2.0, 4.0, 5.8)
APPENDIX_SIZE = 9


@dataclass(frozen=True)
class _Line:
    text: str
    bold: bool = False
    italic: bool = False
    bullet: bool = False


_Cells = tuple[list[_Line], list[_Line]]
"""What a block contributes to the procedure table: teacher cell lines, student cell lines."""


# --- document plumbing -------------------------------------------------------------------------


def _set_font(style: ParagraphStyle, size: int | None = None) -> None:
    """Arial for every script, so Vietnamese diacritics and English share one font."""
    style.font.name = FONT
    if size is not None:
        style.font.size = Pt(size)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), FONT)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)


def _new_document(lesson: Lesson) -> DocxDocument:
    document = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        _set_font(document.styles[name], FONT_SIZE if name == "Normal" else None)
    document.styles["Normal"].paragraph_format.space_after = Pt(4)

    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)

    properties = document.core_properties
    properties.title = f"{lesson_label(lesson)} — {unit_line(lesson)}"
    properties.author = "ClassDeck"
    properties.last_modified_by = "ClassDeck"
    properties.created = lesson.created_at
    properties.modified = lesson.updated_at
    return document


def _shade(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_as_header(table: Table) -> None:
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def _new_table(header: Sequence[str], widths: Sequence[float], document: DocxDocument) -> Table:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    for cell, text, width in zip(table.rows[0].cells, header, widths, strict=True):
        cell.width = Cm(width)
        _shade(cell, HEADER_FILL)
        cell.paragraphs[0].add_run(text).bold = True
    _repeat_as_header(table)
    return table


def _add_row(table: Table, widths: Sequence[float]) -> list[_Cell]:
    cells = table.add_row().cells
    for cell, width in zip(cells, widths, strict=True):
        cell.width = Cm(width)
    return list(cells)


def _write_cell(cell: _Cell, lines: Sequence[_Line]) -> None:
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        if line.bullet:
            paragraph.style = "List Bullet"
        run = paragraph.add_run(line.text)
        run.bold = line.bold or None
        run.italic = line.italic or None


def _mark(paragraph: Paragraph, number: int, *, generated: bool) -> None:
    """Append the AI-added badge (for generated blocks), then the superscript ``[n]`` marker."""
    if generated:
        paragraph.add_run(" ")
        badge = paragraph.add_run(AI_ADDED_BADGE)
        badge.bold = True
        badge.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
    paragraph.add_run(f"[{number}]").font.superscript = True


def _words(value: Enum) -> str:
    return str(value.value).replace("_", " ")


# --- what each block kind says to teachers and pupils ------------------------------------------


def _numbered(steps: Sequence[str]) -> list[_Line]:
    return [_Line(f"{number}. {step}") for number, step in enumerate(steps, start=1)]


def _vocabulary_line(item: VocabItem) -> str:
    text = item.word
    if item.ipa:
        text += f" {item.ipa}" if item.ipa.startswith(("/", "[")) else f" /{item.ipa}/"
    if item.meaning_vi:
        text += f" — {item.meaning_vi}"
    return text


def _activity_cells(block: ActivityBlock) -> _Cells:
    content = block.content
    heading = f"{content.title} ({content.minutes} min, {_words(content.interaction)})"
    teacher = [_Line(heading, bold=True)]
    if content.instructions:
        teacher.append(_Line(f"Instructions: {content.instructions}"))
    teacher.extend(_numbered(content.teacher_steps))
    return teacher, _numbered(content.student_steps)


def _vocabulary_cells(block: VocabularyBlock) -> _Cells:
    teacher = [_Line("Present vocabulary:", bold=True)]
    teacher.extend(_Line(_vocabulary_line(item)) for item in block.content.items)
    return teacher, [_Line("Listen, repeat, point")]


def _pattern_cells(block: SentencePatternBlock) -> _Cells:
    content = block.content
    teacher = [_Line(f"Sentence pattern: {content.pattern}", bold=True)]
    if content.function:
        teacher.append(_Line(content.function, italic=True))
    teacher.extend(_Line(f"Example: {example}") for example in content.examples)
    if content.notes:
        teacher.append(_Line(f"Note: {content.notes}", italic=True))
    return teacher, [_Line("Repeat and practise the pattern")]


def _dialogue_cells(block: DialogueBlock) -> _Cells:
    content = block.content
    teacher = [_Line(content.context, italic=True)] if content.context else []
    teacher.extend(
        _Line(f"{line.speaker}: {line.text}" if line.speaker else line.text)
        for line in content.lines
    )
    return teacher, [_Line("Listen and repeat")]


def _phonics_cells(block: PhonicsBlock) -> _Cells:
    content = block.content
    teacher = [_Line(f"Sound: {content.sound} (letters: {content.letters})", bold=True)]
    if content.words:
        teacher.append(_Line("Words: " + ", ".join(content.words)))
    if content.chant_lines:
        teacher.append(_Line("Chant:", bold=True))
        teacher.extend(_Line(line) for line in content.chant_lines)
    return teacher, [_Line("Listen and repeat the sound")]


def _summary_cells(block: SummaryBlock) -> _Cells:
    teacher = [_Line(bullet, bullet=True) for bullet in block.content.bullets]
    return teacher, [_Line("Recall with the teacher")]


def _exercise_cells(block: Block) -> _Cells:
    """The exercise itself for the students, with its answer key for the teacher."""
    teacher = [_Line(f"{exercise_title(block)} (exercise)", bold=True)]
    instruction = instruction_of(block)
    if instruction:
        teacher.append(_Line(f"Instructions: {instruction}"))
    key = answer_key(block)
    if key:
        teacher.append(_Line("Answer key:", bold=True))
        teacher.extend(_Line(line) for line in key)
    glossary = getattr(block.content, "glossary", {})
    if glossary:
        teacher.append(_Line("Glossary:", bold=True))
        teacher.extend(_Line(f"{word}: {meaning}") for word, meaning in glossary.items())
    students = [_Line(instruction, italic=True)] if instruction else []
    students.extend(_Line(line) for line in exercise_lines(block))
    return teacher, students


def _cells_for(block: Block) -> _Cells | None:
    """The two cells of the block's procedure row, or ``None`` if it has no row."""
    match block:
        case ActivityBlock():
            return _activity_cells(block)
        case VocabularyBlock():
            return _vocabulary_cells(block)
        case SentencePatternBlock():
            return _pattern_cells(block)
        case DialogueBlock():
            return _dialogue_cells(block)
        case PhonicsBlock():
            return _phonics_cells(block)
        case SummaryBlock():
            return _summary_cells(block)
        case _ if block.kind in EXERCISE_KINDS:
            return _exercise_cells(block)
        case _:
            return None


def _objective_text(block: ObjectiveBlock) -> str:
    ids = block.content.outcome_ids
    return f"{block.content.text} ({', '.join(ids)})" if ids else block.content.text


def _teaching_aids(lesson: Lesson) -> list[str]:
    materials = {
        material.strip()
        for block in lesson.ordered_blocks()
        if isinstance(block, ActivityBlock)
        for material in block.content.materials
        if material.strip()
    }
    aids = sorted(materials, key=str.casefold)
    kinds = {block.kind for block in lesson.ordered_blocks()}
    if kinds & {BlockKind.IMAGE_REF, BlockKind.VOCABULARY}:
        aids.append("slides")
    return aids


# --- writer ------------------------------------------------------------------------------------


class _DocxBuilder:
    def __init__(
        self,
        lesson: Lesson,
        sources: Mapping[str, SourceDocument] | None,
        curriculum: Curriculum | None,
    ) -> None:
        self.lesson = lesson
        self.curriculum = curriculum
        self.source_map = build_source_map(lesson, sources)
        self.numbers = {row.block_id: number for number, row in enumerate(self.source_map, 1)}
        self.document = _new_document(lesson)

    def build(self) -> DocxDocument:
        self._header()
        self._objectives()
        self._teaching_aids()
        self._procedure()
        self._wrap_up()
        self._notes()
        self._appendix()
        return self.document

    def _bullets(self, lines: Sequence[str]) -> None:
        for line in lines:
            self.document.add_paragraph(line, style="List Bullet")

    def _header(self) -> None:
        lesson, ref = self.lesson, self.lesson.curriculum
        self.document.add_heading(unit_line(lesson), level=0)

        number, name = unit_number(ref), unit_title(lesson)
        unit = " · ".join(str(part) for part in (number, name) if part)
        fields = [
            ("Unit", unit or "______"),
            ("Grade", str(lesson.grade)),
            ("Lesson", str(ref.lesson_no) if ref.lesson_no is not None else "______"),
            self._period_field(),
            ("Date", "______"),
            ("Class", "______"),
        ]
        table = self.document.add_table(rows=3, cols=2)
        table.style = "Table Grid"
        for index, (label, value) in enumerate(fields):
            paragraph = table.cell(index // 2, index % 2).paragraphs[0]
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(value)

    def _period_field(self) -> tuple[str, str]:
        """``Period: 2`` for a primary lesson; a secondary lesson is one period: its length."""
        lesson, ref = self.lesson, self.lesson.curriculum
        if lesson.grade > 5 and ref.period is None:
            return "Time", f"{sum(section.minutes for section in lesson.sections)} min"
        return "Period", str(ref.period) if ref.period is not None else "______"

    def _objectives(self) -> None:
        self.document.add_heading("I. Objectives", level=1)
        objectives = [
            block for block in self.lesson.ordered_blocks() if isinstance(block, ObjectiveBlock)
        ]
        self._bullets([_objective_text(block) for block in objectives])
        outcome_ids = self.lesson.curriculum.outcome_ids
        if self.curriculum is None or not outcome_ids:
            return
        self.document.add_paragraph().add_run("Curriculum outcomes:").bold = True
        lines = []
        for outcome_id in outcome_ids:
            try:
                text = self.curriculum.outcome(outcome_id).text
            except KeyError:
                text = "(not found in the curriculum)"
            lines.append(f"{outcome_id} — {text}")
        self._bullets(lines)

    def _teaching_aids(self) -> None:
        self.document.add_heading("II. Teaching aids", level=1)
        self._bullets(_teaching_aids(self.lesson) or ["None specified."])

    def _procedure(self) -> None:
        self.document.add_heading("III. Procedure", level=1)
        for section in self.lesson.sections:
            if section.phase is not Phase.WRAP_UP:
                self._section(section)

    def _wrap_up(self) -> None:
        self.document.add_heading("IV. Wrap-up / Homework", level=1)
        shown = False
        for section in self.lesson.sections:
            if section.phase is Phase.WRAP_UP:
                shown = self._section(section) or shown
        if not shown:
            self.document.add_paragraph(WRAP_UP_FALLBACK)

    def _section(self, section: Section) -> bool:
        """Write one phase as a heading, a two-column table and its notes; True if any content."""
        blocks = [self.lesson.blocks[block_id] for block_id in section.block_ids]
        rows = [(block, cells) for block in blocks if (cells := _cells_for(block)) is not None]
        notes = [
            (text, block)
            for block in blocks
            for text in (
                block.content.text if isinstance(block, TeacherNoteBlock) else None,
                block.teacher_note,
            )
            if text
        ]
        self.document.add_heading(f"{section.title} ({section.minutes} min)", level=2)
        if rows:
            self._procedure_table(rows)
        for text, block in notes:
            paragraph = self.document.add_paragraph()
            paragraph.add_run(text).italic = True
            _mark(
                paragraph,
                self.numbers[block.id],
                generated=block.provenance is Provenance.GENERATED,
            )
        return bool(rows or notes)

    def _procedure_table(self, rows: Sequence[tuple[Block, _Cells]]) -> None:
        table = _new_table(
            ("Teacher's activities", "Students' activities"), PROCEDURE_WIDTHS, self.document
        )
        for block, (teacher, students) in rows:
            teacher_cell, students_cell = _add_row(table, PROCEDURE_WIDTHS)
            _write_cell(teacher_cell, teacher)
            _write_cell(students_cell, students)
            _mark(
                teacher_cell.paragraphs[-1],
                self.numbers[block.id],
                generated=block.provenance is Provenance.GENERATED,
            )
            if isinstance(block, ActivityBlock) and block.content.competency_tags:
                self._competences(table, block)

    @staticmethod
    def _competences(table: Table, block: ActivityBlock) -> None:
        cells = _add_row(table, PROCEDURE_WIDTHS)
        merged = cells[0].merge(cells[1])
        run = merged.paragraphs[0].add_run(
            "Competences: " + ", ".join(block.content.competency_tags)
        )
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _notes(self) -> None:
        warnings = collect_warnings(self.lesson)
        if not warnings:
            return
        self.document.add_heading("Notes", level=1)
        self._bullets([format_check(block_id, check) for block_id, check in warnings])

    def _appendix(self) -> None:
        self.document.add_heading("Appendix — Source map", level=1)
        table = _new_table(
            ("n", "Block", "Kind", "Provenance", "Citation", "Quote"),
            APPENDIX_WIDTHS,
            self.document,
        )
        for number, row in enumerate(self.source_map, start=1):
            cells = _add_row(table, APPENDIX_WIDTHS)
            values = (
                str(number),
                row.block_id,
                _words(row.kind),
                PROVENANCE_LABEL[row.provenance],
                origin_text(row),
                row.quote or "",
            )
            for cell, value in zip(cells, values, strict=True):
                cell.paragraphs[0].add_run(value)
        for table_row in table.rows:
            for cell in table_row.cells:
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(APPENDIX_SIZE)


def export_docx(
    lesson: Lesson,
    out_path: Path,
    *,
    sources: Mapping[str, SourceDocument] | None = None,
    curriculum: Curriculum | None = None,
) -> Path:
    """Write the lesson plan and return the path.

    Raises :class:`~classdeck.export.source_map.ExportBlocked` if any check failed. With
    ``curriculum`` the lesson's outcome ids are expanded to their text; with ``sources`` the
    source map quotes the cited spans for citations that carry no quote.
    """
    ensure_exportable(lesson)
    document = _DocxBuilder(lesson, sources, curriculum).build()
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
