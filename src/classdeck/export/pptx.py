"""PowerPoint writer: the slides for one lesson, each labelled with where its content came from.

Slide shapes follow the export conventions in ``docs/english-rules.md``: flashcards for
vocabulary, a pattern slide with the slot highlighted, speaker labels on dialogues. Every content
slide carries a footer with the block id and its source; ``generated`` blocks also carry an
AI-added badge and a speaker note. Teacher-only material (steps, materials, notes) goes into the
speaker notes. A source map closes the deck.
"""

from __future__ import annotations

import math
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as Deck
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import _Paragraph
from pptx.text.text import _Run as PptxRun
from pptx.util import Emu, Inches, Pt

from classdeck.export.source_map import (
    AI_ADDED_BADGE,
    AI_ADDED_NOTE,
    PROVENANCE_LABEL,
    build_source_map,
    collect_warnings,
    ensure_exportable,
    format_check,
    format_citation,
    origin_text,
)
from classdeck.lesson import (
    ActivityBlock,
    Block,
    Citation,
    DialogueBlock,
    ImageRefBlock,
    Lesson,
    ObjectiveBlock,
    PhonicsBlock,
    Provenance,
    Section,
    SentencePatternBlock,
    SourceDocument,
    SpanKind,
    SummaryBlock,
    TeacherNoteBlock,
    VocabularyBlock,
)

FONT = "Arial"
MIN_SIZE = 18

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
ACCENT = RGBColor(0x0B, 0x6E, 0x99)
BADGE_FILL = RGBColor(0xB4, 0x3C, 0x0A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_WIDTH = Emu(12192000)  # 13.333 in
SLIDE_HEIGHT = Emu(6858000)  # 7.5 in

Box = tuple[float, float, float, float]
"""Left, top, width, height in inches."""

MARGIN = 0.6
CONTENT_WIDTH = 12.13
TITLE_BOX: Box = (MARGIN, 0.35, 10.2, 1.0)
BODY_TOP = 1.5
FOOTER_BOX: Box = (MARGIN, 6.8, CONTENT_WIDTH, 0.5)
BADGE_BOX: Box = (11.03, 0.45, 1.7, 0.5)

DIALOGUE_LINES_PER_SLIDE = 6
NOTES_PER_SLIDE = 8
MAX_TABLE_ROWS = 12
TABLE_QUOTE_LIMIT = 60

_SLOT = re.compile(r"(_{2,})")


# --- text helpers ------------------------------------------------------------------------------


@dataclass(frozen=True)
class _TextRun:
    text: str
    bold: bool = False
    color: RGBColor | None = None


@dataclass(frozen=True)
class _Para:
    runs: tuple[_TextRun, ...]
    size: int
    color: RGBColor = INK
    italic: bool = False
    bullet: bool = False
    space_after: int = 6
    align: PP_ALIGN | None = None


def _para(
    text: str,
    size: int,
    *,
    bold: bool = False,
    color: RGBColor = INK,
    italic: bool = False,
    bullet: bool = False,
    space_after: int = 6,
    align: PP_ALIGN | None = None,
) -> _Para:
    return _Para((_TextRun(text, bold),), size, color, italic, bullet, space_after, align)


def _style_run(run: PptxRun, size: int, bold: bool, color: RGBColor, italic: bool) -> None:
    font = run.font
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color


def _bullet(paragraph: _Paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(Inches(0.4)))
    properties.set("indent", str(-Inches(0.4)))
    marker = OxmlElement("a:buChar")
    marker.set("char", "•")
    properties.append(marker)


def _add_rich(
    slide: Slide,
    box: Box,
    paragraphs: Sequence[_Para],
    *,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    name: str | None = None,
) -> Shape:
    left, top, width, height = box
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    for index, para in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = para.align or align
        paragraph.space_after = Pt(para.space_after)
        for run in para.runs:
            added = paragraph.add_run()
            added.text = run.text
            _style_run(added, para.size, run.bold, run.color or para.color, para.italic)
        if para.bullet:
            _bullet(paragraph)
    if name:
        shape.name = name
    return shape


def _add_text(
    slide: Slide,
    box: Box,
    content: str | Sequence[str],
    *,
    size: int,
    bold: bool = False,
    color: RGBColor = INK,
    italic: bool = False,
    bullets: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    space_after: int = 6,
    name: str | None = None,
) -> Shape:
    """A word-wrapped textbox in one style; a sequence of strings becomes one paragraph each."""
    lines = [content] if isinstance(content, str) else list(content)
    paragraphs = [
        _para(
            line,
            size,
            bold=bold,
            color=color,
            italic=italic,
            bullet=bullets,
            space_after=space_after,
        )
        for line in lines
    ]
    return _add_rich(slide, box, paragraphs, align=align, anchor=anchor, name=name)


def _fit_size(texts: Sequence[str], base: int) -> int:
    """Shrink the font as the text grows, never below the readable minimum."""
    lines = sum(max(1, math.ceil(len(text) / 60)) for text in texts)
    return max(MIN_SIZE, base - 2 * max(0, lines - 6))


def _title_size(title: str) -> int:
    if len(title) <= 34:
        return 32
    return 28 if len(title) <= 50 else 24


def _chunks[T](items: Sequence[T], size: int) -> Iterator[Sequence[T]]:
    for start in range(0, len(items), size):
        yield items[start : start + size]


def _words(value: Enum) -> str:
    return str(value.value).replace("_", " ")


def _slot_runs(pattern: str) -> tuple[_TextRun, ...]:
    """Split a pattern so each ``___`` slot becomes a bold run in the accent colour."""
    return tuple(
        _TextRun(part, bold=True, color=ACCENT) if part.startswith("__") else _TextRun(part)
        for part in _SLOT.split(pattern)
        if part
    )


# --- labelling ---------------------------------------------------------------------------------


def _derived_refs(block: Block) -> list[str]:
    refs = list(block.depends_on)
    if isinstance(block, ObjectiveBlock):
        refs.extend(block.content.outcome_ids)
    return refs


def _footer_text(blocks: Sequence[Block]) -> str:
    """Block id(s), then ``Source: p. 6, ¶3`` for cited blocks and ``Derived from …``."""
    parts = [", ".join(block.id for block in blocks)]
    citations = list(dict.fromkeys(format_citation(c) for block in blocks for c in block.sources))
    if citations:
        parts.append("Source: " + ", ".join(citations))
    derived = [block for block in blocks if block.provenance is Provenance.DERIVED]
    if derived:
        refs = list(dict.fromkeys(ref for block in derived for ref in _derived_refs(block)))
        parts.append("Derived from " + (", ".join(refs) or "curriculum"))
    return " · ".join(parts)


def _append_notes(slide: Slide, lines: Sequence[str]) -> None:
    lines = [line for line in lines if line]
    if not lines:
        return
    frame = slide.notes_slide.notes_text_frame
    if frame is None:  # pragma: no cover - the stock notes master always has a body
        return
    for line in lines:
        if len(frame.paragraphs) == 1 and not frame.paragraphs[0].text:
            frame.paragraphs[0].text = line
        else:
            frame.add_paragraph().text = line


def _add_picture(slide: Slide, path: Path, box: Box, *, alt: str) -> None:
    """Place an image inside ``box``, scaled to fit and centred, with its alt text set."""
    left, top, width, height = (Inches(value) for value in box)
    picture = slide.shapes.add_picture(str(path), left, top)
    scale = min(width / picture.width, height / picture.height)
    fitted_width, fitted_height = int(picture.width * scale), int(picture.height * scale)
    picture.width, picture.height = Emu(fitted_width), Emu(fitted_height)
    picture.left = Emu(left + (width - fitted_width) // 2)
    picture.top = Emu(top + (height - fitted_height) // 2)
    picture.name = "Picture"
    picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", alt)


def _add_table(
    slide: Slide,
    header: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float],
    *,
    name: str,
    size: int = 12,
    row_height: float = 0.4,
    top: float = BODY_TOP,
) -> None:
    count = len(rows) + 1
    shape = slide.shapes.add_table(
        count,
        len(header),
        Inches(MARGIN),
        Inches(top),
        Inches(sum(widths)),
        Inches(row_height * count),
    )
    shape.name = name
    table = shape.table
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, values in enumerate([header, *rows]):
        table.rows[row_index].height = Inches(row_height)
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text_frame.word_wrap = True
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = value
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT
                _style_run(run, size, True, WHITE, False)
            else:
                _style_run(run, size, False, INK, False)


# --- the deck ----------------------------------------------------------------------------------


class _DeckBuilder:
    def __init__(
        self,
        lesson: Lesson,
        sources: Mapping[str, SourceDocument] | None,
        image_root: Path | None,
    ) -> None:
        self.lesson = lesson
        self.sources = sources
        self.image_root = image_root
        self.deck: Deck = Presentation()
        self.deck.slide_width = SLIDE_WIDTH
        self.deck.slide_height = SLIDE_HEIGHT
        self._layout = self.deck.slide_layouts.get_by_name("Title Only")
        self._last: Slide | None = None

    def build(self) -> Deck:
        self._title_slide()
        for section in self.lesson.sections:
            self._section_slide(section)
            self._blocks([self.lesson.blocks[block_id] for block_id in section.block_ids])
        self._notes_slides()
        self._source_map_slides()
        self._set_properties()
        return self.deck

    # -- plumbing

    def _slide(
        self,
        title: str,
        *,
        box: Box = TITLE_BOX,
        size: int | None = None,
        align: PP_ALIGN = PP_ALIGN.LEFT,
        anchor: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    ) -> Slide:
        slide = self.deck.slides.add_slide(self._layout)
        heading = slide.shapes.title
        if heading is None:  # pragma: no cover - the stock layout always has a title
            raise RuntimeError("slide layout has no title placeholder")
        left, top, width, height = box
        heading.left, heading.top = Inches(left), Inches(top)
        heading.width, heading.height = Inches(width), Inches(height)
        frame = heading.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = title
        _style_run(run, size or _title_size(title), True, ACCENT, False)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.15))
        bar.name = "Accent bar"
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        self._last = slide
        return slide

    def _content_slide(
        self,
        title: str,
        blocks: Sequence[Block],
        *,
        first: bool = True,
        notes: Sequence[str] = (),
    ) -> Slide:
        """A titled slide with the footer, the AI-added badge and the standard speaker notes."""
        slide = self._slide(title)
        _add_text(slide, FOOTER_BOX, _footer_text(blocks), size=18, color=MUTED, name="Footer")
        lines: list[str] = []
        if any(block.provenance is Provenance.GENERATED for block in blocks):
            self._badge(slide)
            lines.append(AI_ADDED_NOTE)
        if first:
            lines.extend(block.teacher_note for block in blocks if block.teacher_note)
        _append_notes(slide, [*lines, *notes])
        return slide

    @staticmethod
    def _badge(slide: Slide) -> None:
        badge = _add_text(
            slide,
            BADGE_BOX,
            AI_ADDED_BADGE,
            size=18,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            name="AI-added badge",
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = BADGE_FILL

    def _resolve_image(self, citation: Citation | None) -> Path | None:
        """The image file behind a citation, only if it is an image span inside ``image_root``."""
        if citation is None or self.sources is None or self.image_root is None:
            return None
        document = self.sources.get(citation.doc_id)
        span = document.find_span(citation.page, citation.span_id) if document else None
        if span is None or span.kind is not SpanKind.IMAGE or not span.image_path:
            return None
        root = self.image_root.resolve()
        path = (root / span.image_path).resolve()
        return path if path.is_relative_to(root) and path.is_file() else None

    # -- framing slides

    def _title_slide(self) -> None:
        ref = self.lesson.curriculum
        parts = [
            f"Grade {self.lesson.grade}",
            ref.unit_id,
            f"Lesson {ref.lesson_no}" if ref.lesson_no is not None else "",
            f"Period {ref.period}" if ref.period is not None else "",
        ]
        slide = self._slide(
            self.lesson.title,
            box=(MARGIN, 2.0, CONTENT_WIDTH, 1.9),
            size=44,
            align=PP_ALIGN.CENTER,
        )
        _add_text(
            slide,
            (MARGIN, 4.1, CONTENT_WIDTH, 0.8),
            " · ".join(part for part in parts if part),
            size=28,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )

    def _section_slide(self, section: Section) -> None:
        self._slide(
            f"{section.title} — {section.minutes} min",
            box=(MARGIN, 2.4, CONTENT_WIDTH, 2.4),
            size=44,
            align=PP_ALIGN.CENTER,
        )

    def _notes_slides(self) -> None:
        warnings = collect_warnings(self.lesson)
        pages = list(_chunks(warnings, NOTES_PER_SLIDE))
        for index, page in enumerate(pages, start=1):
            title = "Notes" if index == 1 else "Notes (continued)"
            slide = self._slide(title)
            lines = [format_check(block_id, check) for block_id, check in page]
            _add_text(
                slide,
                (MARGIN, BODY_TOP, CONTENT_WIDTH, 5.2),
                lines,
                size=_fit_size(lines, 22),
                bullets=True,
                space_after=10,
            )

    def _source_map_slides(self) -> None:
        rows = build_source_map(self.lesson, self.sources)
        pages = list(_chunks(rows, MAX_TABLE_ROWS)) or [[]]
        for index, page in enumerate(pages, start=1):
            title = "Source map" if len(pages) == 1 else f"Source map ({index}/{len(pages)})"
            slide = self._slide(title)
            _add_table(
                slide,
                ["Block", "Kind", "Provenance", "Citation", "Quote"],
                [
                    [
                        row.block_id,
                        _words(row.kind),
                        PROVENANCE_LABEL[row.provenance],
                        origin_text(row),
                        _shorten(row.quote or "", TABLE_QUOTE_LIMIT),
                    ]
                    for row in page
                ],
                [1.0, 1.9, 1.4, 2.4, 5.43],
                name="Source map table",
            )

    def _set_properties(self) -> None:
        properties = self.deck.core_properties
        properties.title = self.lesson.title
        properties.author = "ClassDeck"
        properties.last_modified_by = "ClassDeck"
        properties.created = self.lesson.created_at
        properties.modified = self.lesson.updated_at

    # -- blocks

    def _blocks(self, blocks: Sequence[Block]) -> None:
        objectives = [block for block in blocks if isinstance(block, ObjectiveBlock)]
        for block in blocks:
            match block:
                case ObjectiveBlock():
                    if block is objectives[0]:
                        self._objectives(objectives)
                case VocabularyBlock():
                    self._vocabulary(block)
                case SentencePatternBlock():
                    self._pattern(block)
                case DialogueBlock():
                    self._dialogue(block)
                case PhonicsBlock():
                    self._phonics(block)
                case ActivityBlock():
                    self._activity(block)
                case SummaryBlock():
                    self._summary(block)
                case ImageRefBlock():
                    self._image_ref(block)
                case TeacherNoteBlock():
                    self._teacher_note(block)

    def _objectives(self, blocks: Sequence[ObjectiveBlock]) -> None:
        texts = [block.content.text for block in blocks]
        slide = self._content_slide("Objectives", blocks)
        _add_text(
            slide,
            (MARGIN, BODY_TOP, CONTENT_WIDTH, 5.1),
            texts,
            size=_fit_size(texts, 30),
            bullets=True,
            space_after=14,
        )

    def _vocabulary(self, block: VocabularyBlock) -> None:
        items = block.content.items
        for index, item in enumerate(items):
            title = "Vocabulary" if len(items) == 1 else f"Vocabulary {index + 1}/{len(items)}"
            notes = [f"Example: {item.example}"] if item.example else []
            if item.image is not None:
                notes.append(f"Picture: {format_citation(item.image)}")
            slide = self._content_slide(title, [block], first=index == 0, notes=notes)

            picture = self._resolve_image(item.image)
            if picture is not None:
                _add_picture(slide, picture, (7.9, BODY_TOP, 4.83, 5.1), alt=item.word)
                box, align = (MARGIN, BODY_TOP, 7.0, 5.1), PP_ALIGN.LEFT
            else:
                box, align = (MARGIN, BODY_TOP, CONTENT_WIDTH, 5.1), PP_ALIGN.CENTER
            paragraphs = [_para(item.word, 66, bold=True, color=ACCENT, space_after=12)]
            if item.ipa:
                paragraphs.append(_para(item.ipa, 36, space_after=12))
            if item.meaning_vi:
                paragraphs.append(_para(item.meaning_vi, 24, color=MUTED))
            _add_rich(slide, box, paragraphs, align=align, anchor=MSO_ANCHOR.MIDDLE)

    def _pattern(self, block: SentencePatternBlock) -> None:
        content = block.content
        slide = self._content_slide(
            "Sentence pattern", [block], notes=[content.notes] if content.notes else []
        )
        pattern = _Para(_slot_runs(content.pattern), 44, space_after=8)
        _add_rich(
            slide,
            (MARGIN, 1.6, CONTENT_WIDTH, 1.6),
            [pattern],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if content.function:
            _add_text(
                slide,
                (MARGIN, 3.2, CONTENT_WIDTH, 0.6),
                content.function,
                size=24,
                color=MUTED,
                italic=True,
                align=PP_ALIGN.CENTER,
            )
        if content.examples:
            _add_text(
                slide,
                (1.4, 4.0, CONTENT_WIDTH - 1.6, 2.6),
                content.examples,
                size=_fit_size(content.examples, 28),
                bullets=True,
                space_after=10,
            )

    def _dialogue(self, block: DialogueBlock) -> None:
        content = block.content
        pages = list(_chunks(content.lines, DIALOGUE_LINES_PER_SLIDE))
        for index, page in enumerate(pages, start=1):
            title = "Dialogue" if len(pages) == 1 else f"Dialogue ({index}/{len(pages)})"
            slide = self._content_slide(title, [block], first=index == 1)
            top = BODY_TOP
            if content.context:
                _add_text(
                    slide,
                    (MARGIN, 1.3, CONTENT_WIDTH, 0.6),
                    content.context,
                    size=22,
                    color=MUTED,
                    italic=True,
                )
                top = 2.0
            size = _fit_size([line.text for line in page], 30)
            paragraphs = [
                _Para(
                    (
                        _TextRun(f"{line.speaker}: ", bold=True, color=ACCENT),
                        _TextRun(line.text),
                    )
                    if line.speaker
                    else (_TextRun(line.text),),
                    size,
                    space_after=14,
                )
                for line in page
            ]
            _add_rich(slide, (MARGIN, top, CONTENT_WIDTH, 6.7 - top), paragraphs)

    def _phonics(self, block: PhonicsBlock) -> None:
        content = block.content
        slide = self._content_slide("Phonics", [block])
        _add_rich(
            slide,
            (MARGIN, BODY_TOP, 4.6, 5.1),
            [
                _para(content.sound, 96, bold=True, color=ACCENT, space_after=8),
                _para(content.letters, 54, bold=True),
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        paragraphs: list[_Para] = []
        if content.words:
            paragraphs.append(_para("Words", 20, bold=True, color=MUTED, space_after=4))
            paragraphs.append(_para("  ·  ".join(content.words), 36, bold=True, space_after=18))
        if content.chant_lines:
            chant_size = _fit_size(content.chant_lines, 26)
            paragraphs.append(_para("Chant", 20, bold=True, color=MUTED, space_after=4))
            paragraphs.extend(
                _para(line, chant_size, italic=True, space_after=8) for line in content.chant_lines
            )
        if paragraphs:
            _add_rich(slide, (5.5, BODY_TOP, 7.23, 5.1), paragraphs, anchor=MSO_ANCHOR.MIDDLE)

    def _activity(self, block: ActivityBlock) -> None:
        content = block.content
        details = [
            f"Activity type: {_words(content.type)} · Interaction: {_words(content.interaction)}"
            f" · {content.minutes} min"
        ]
        if content.skills:
            details.append("Skills: " + ", ".join(skill.value for skill in content.skills))
        if content.materials:
            details.append("Materials: " + ", ".join(content.materials))
        if content.teacher_steps:
            details.append("Teacher steps:")
            details.extend(f"{n}. {step}" for n, step in enumerate(content.teacher_steps, start=1))
        slide = self._content_slide(content.title, [block], notes=details)

        texts = [content.instructions, *content.student_steps]
        size = _fit_size([text for text in texts if text], 28)
        paragraphs: list[_Para] = []
        if content.instructions:
            paragraphs.append(_para(content.instructions, size, space_after=16))
        paragraphs.extend(
            _para(step, size - 2, bullet=True, space_after=10) for step in content.student_steps
        )
        if paragraphs:
            _add_rich(slide, (MARGIN, BODY_TOP, CONTENT_WIDTH, 5.1), paragraphs)

    def _summary(self, block: SummaryBlock) -> None:
        bullets = block.content.bullets
        slide = self._content_slide("Summary", [block])
        _add_text(
            slide,
            (MARGIN, BODY_TOP, CONTENT_WIDTH, 5.1),
            bullets,
            size=_fit_size(bullets, 30),
            bullets=True,
            space_after=14,
        )

    def _image_ref(self, block: ImageRefBlock) -> None:
        content = block.content
        picture = next(
            (path for c in block.sources if (path := self._resolve_image(c)) is not None), None
        )
        slide = self._content_slide("Picture", [block], notes=[f"Picture: {content.alt}"])
        if picture is not None:
            _add_picture(slide, picture, (MARGIN, BODY_TOP, CONTENT_WIDTH, 4.3), alt=content.alt)
        else:
            _add_text(
                slide,
                (MARGIN, BODY_TOP, CONTENT_WIDTH, 4.3),
                f"Picture not available: {content.alt}",
                size=24,
                color=MUTED,
                italic=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        if content.caption:
            _add_text(
                slide,
                (MARGIN, 5.9, CONTENT_WIDTH, 0.8),
                content.caption,
                size=26,
                align=PP_ALIGN.CENTER,
            )

    def _teacher_note(self, block: TeacherNoteBlock) -> None:
        """Teacher notes get no slide of their own; they join the previous slide's notes."""
        label = "Teacher note"
        if block.provenance is Provenance.GENERATED:
            label += f" ({AI_ADDED_BADGE})"
        if self._last is not None:
            _append_notes(self._last, [f"{label}: {block.content.text}", block.teacher_note or ""])


def _shorten(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def export_pptx(
    lesson: Lesson,
    out_path: Path,
    *,
    sources: Mapping[str, SourceDocument] | None = None,
    image_root: Path | None = None,
) -> Path:
    """Write the lesson as a 16:9 deck and return the path.

    Raises :class:`~classdeck.export.source_map.ExportBlocked` if any check failed. Pictures
    are placed only when ``sources`` and ``image_root`` are given and the cited image span's
    file exists under ``image_root``.
    """
    ensure_exportable(lesson)
    deck = _DeckBuilder(lesson, sources, image_root).build()
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))
    return path
