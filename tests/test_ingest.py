"""Ingest: spans, ids, images and stability for PDF, DOCX and PPTX."""

from __future__ import annotations

import shutil
from collections.abc import Callable
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from classdeck.ingest import (
    EXTRACTION_VERSION,
    UnsupportedFormat,
    ingest,
    sha256_of,
)
from classdeck.lesson import SourceDocument, SourceStore, Span, SpanKind

FIXTURES = Path(__file__).resolve().parent.parent / "eval" / "fixtures" / "public"

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MEDIA_TYPES = {"pdf": "application/pdf", "docx": DOCX_TYPE, "pptx": PPTX_TYPE}
HAS_BBOX = {"pdf": True, "docx": False, "pptx": True}

NUMBERED_LINES = [
    "1. Look, listen and repeat.",
    "2. Listen, point and say.",
    "3. Let's talk.",
    "4. Listen and tick.",
    "5. Look, complete and read.",
    "6. Let's sing.",
]
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


@pytest.fixture(params=["pdf", "docx", "pptx"])
def fmt(request: pytest.FixtureRequest) -> str:
    return request.param


@pytest.fixture
def fixture_path(fmt: str) -> Path:
    return FIXTURES / f"mini-unit-hello.{fmt}"


def all_spans(document: SourceDocument) -> list[Span]:
    return [span for page in document.pages for span in page.spans]


def kinds_and_text(page_spans: list[Span]) -> list[tuple[str, str]]:
    return [(span.kind.value, span.text) for span in page_spans]


# --- the public fixtures, all three formats ----------------------------------------------------


def test_two_pages(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    assert [page.number for page in document.pages] == [1, 2]


def test_headings(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    headings = [span for span in all_spans(document) if span.kind is SpanKind.HEADING]
    assert [span.span_id for span in headings] == ["p1.s1", "p2.s1"]
    assert all("Unit 1" in span.text for span in headings)
    assert "Period 2" in headings[1].text


def test_numbered_lines_are_list_items(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    by_text = {span.text: span for span in all_spans(document)}
    for line in NUMBERED_LINES:
        assert by_text[line].kind is SpanKind.LIST_ITEM, line
    list_items = [s for s in all_spans(document) if s.kind is SpanKind.LIST_ITEM]
    assert [span.text for span in list_items] == NUMBERED_LINES


def test_dialogue_lines_are_paragraphs(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    paragraphs = [s.text for s in all_spans(document) if s.kind is SpanKind.PARAGRAPH]
    assert "Hello, I'm Ha." in paragraphs
    assert "Hi, Ha. I'm Tien." in paragraphs
    assert "Hello. / Hi. / I'm (name)." in paragraphs
    assert len(paragraphs) == 8


def test_page_content_in_reading_order(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    assert kinds_and_text(document.pages[0].spans) == [
        ("heading", "Unit 1 Hello"),
        ("list_item", "1. Look, listen and repeat."),
        ("image", ""),
        ("paragraph", "Hello, I'm Ha."),
        ("paragraph", "Hi, Ha. I'm Tien."),
        ("list_item", "2. Listen, point and say."),
        ("paragraph", "Hello, I'm ___."),
        ("paragraph", "Hi, ___. I'm ___."),
        ("list_item", "3. Let's talk."),
        ("paragraph", "Hello. / Hi. / I'm (name)."),
    ]
    assert kinds_and_text(document.pages[1].spans) == [
        ("heading", "Unit 1 Hello (Lesson 1, Period 2)"),
        ("list_item", "4. Listen and tick."),
        ("paragraph", "a. Hello, I'm Duc. b. Hi, I'm Lan."),
        ("list_item", "5. Look, complete and read."),
        ("paragraph", "Hello, I'm _____. Hi, _____. I'm _____."),
        ("list_item", "6. Let's sing."),
        (
            "paragraph",
            "Hello, hello, hello! I'm Ha. Hello, hello, hello! I'm Tien.",
        ),
    ]


def test_image_span_and_file(fixture_path: Path, tmp_path: Path) -> None:
    document = ingest(fixture_path, image_dir=tmp_path)
    images = [span for span in all_spans(document) if span.kind is SpanKind.IMAGE]
    assert len(images) == 1
    image = images[0]
    assert image.text == ""
    assert image.span_id.startswith("p1.")
    assert image.image_path == f"{document.doc_id}/p1_img1.png"
    assert not Path(image.image_path).is_absolute()
    written = tmp_path / image.image_path
    assert written.read_bytes().startswith(PNG_MAGIC)


def test_image_spans_without_image_dir(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    images = [span for span in all_spans(document) if span.kind is SpanKind.IMAGE]
    assert len(images) == 1
    assert images[0].image_path is None
    with_dir_ids = [s.span_id for s in all_spans(document)]
    assert with_dir_ids == [s.span_id for s in all_spans(ingest(fixture_path))]


def test_bbox(fmt: str, fixture_path: Path) -> None:
    document = ingest(fixture_path)
    for span in all_spans(document):
        if HAS_BBOX[fmt]:
            assert span.bbox is not None, span.span_id
            x0, y0, x1, y1 = span.bbox
            assert x1 > x0
            assert y1 > y0
        else:
            assert span.bbox is None


def test_span_ids_are_consecutive_per_page(fixture_path: Path) -> None:
    document = ingest(fixture_path)
    for page in document.pages:
        assert [span.order for span in page.spans] == list(range(1, len(page.spans) + 1))
        assert [span.span_id for span in page.spans] == [
            f"p{page.number}.s{order}" for order in range(1, len(page.spans) + 1)
        ]


def test_ingest_is_stable(fixture_path: Path, tmp_path: Path) -> None:
    first = ingest(fixture_path, image_dir=tmp_path / "a")
    second = ingest(fixture_path, image_dir=tmp_path / "a")
    assert first == second
    third = ingest(fixture_path)
    assert [(s.span_id, s.kind, s.text) for s in all_spans(first)] == [
        (s.span_id, s.kind, s.text) for s in all_spans(third)
    ]


def test_document_metadata(fmt: str, fixture_path: Path) -> None:
    document = ingest(fixture_path)
    digest = sha256_of(fixture_path)
    assert document.sha256 == digest
    assert document.doc_id == "d_" + digest[:8]
    assert document.filename == fixture_path.name
    assert document.media_type == MEDIA_TYPES[fmt]
    assert document.extraction_version == EXTRACTION_VERSION


def test_formats_agree_on_text() -> None:
    texts = {
        fmt: [
            (page.number, kind, text)
            for page in ingest(FIXTURES / f"mini-unit-hello.{fmt}").pages
            for kind, text in kinds_and_text(page.spans)
        ]
        for fmt in MEDIA_TYPES
    }
    assert texts["pdf"] == texts["docx"] == texts["pptx"]


def test_round_trip_through_source_store(fixture_path: Path, tmp_path: Path) -> None:
    document = ingest(fixture_path, image_dir=tmp_path / "images")
    store = SourceStore(tmp_path / "store")
    store.save(document)
    loaded = store.load(document.doc_id)
    assert loaded == document
    span = loaded.find_span(1, "p1.s2")
    assert span is not None
    assert (span.kind, span.text) == (SpanKind.LIST_ITEM, "1. Look, listen and repeat.")
    assert loaded.find_span(1, "p9.s9") is None


# --- dispatch ----------------------------------------------------------------------------------


def test_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Hello", encoding="utf-8")
    with pytest.raises(UnsupportedFormat):
        ingest(path)


def test_suffix_is_case_insensitive(tmp_path: Path) -> None:
    copy = tmp_path / "UNIT.PDF"
    shutil.copy(FIXTURES / "mini-unit-hello.pdf", copy)
    document = ingest(copy)
    assert document.media_type == MEDIA_TYPES["pdf"]
    assert document.filename == "UNIT.PDF"


def test_missing_file_is_reported(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        ingest(tmp_path / "missing.pdf")


# --- PDF details -------------------------------------------------------------------------------


def make_pdf(path: Path, draw: Callable[[pymupdf.Page], None]) -> Path:
    document = pymupdf.open()
    draw(document.new_page())
    document.save(str(path))
    document.close()
    return path


def test_pdf_tight_lines_still_split_at_numbered_items(tmp_path: Path) -> None:
    def draw(page: pymupdf.Page) -> None:
        page.insert_text((60, 80), "Unit 2  Big", fontsize=22)
        for index, line in enumerate(NUMBERED_LINES[:3]):
            page.insert_text((60, 110 + 15 * index), line, fontsize=12)

    document = ingest(make_pdf(tmp_path / "tight.pdf", draw))
    assert kinds_and_text(document.pages[0].spans) == [
        ("heading", "Unit 2 Big"),
        *[("list_item", line) for line in NUMBERED_LINES[:3]],
    ]


def test_pdf_small_images_are_not_spans(tmp_path: Path) -> None:
    def draw(page: pymupdf.Page) -> None:
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
        pixmap.clear_with(120)
        page.insert_text((60, 80), "Bullets and a picture", fontsize=12)
        page.insert_image(pymupdf.Rect(60, 100, 70, 110), pixmap=pixmap)
        page.insert_image(pymupdf.Rect(60, 200, 160, 260), pixmap=pixmap)

    document = ingest(make_pdf(tmp_path / "images.pdf", draw), image_dir=tmp_path / "out")
    images = [span for span in all_spans(document) if span.kind is SpanKind.IMAGE]
    assert len(images) == 1
    assert images[0].bbox == (80.0, 200.0, 140.0, 260.0)
    assert images[0].image_path is not None
    assert (tmp_path / "out" / images[0].image_path).is_file()


def test_pdf_table_replaces_its_cell_text(tmp_path: Path) -> None:
    def draw(page: pymupdf.Page) -> None:
        page.insert_text((60, 80), "Words", fontsize=12)
        for row in range(2):
            for column in range(2):
                cell = pymupdf.Rect(
                    60 + 100 * column, 200 + 30 * row, 160 + 100 * column, 230 + 30 * row
                )
                page.draw_rect(cell, color=(0, 0, 0), width=1)
                page.insert_text((cell.x0 + 8, cell.y0 + 20), "ABCD"[2 * row + column], fontsize=12)

    document = ingest(make_pdf(tmp_path / "table.pdf", draw))
    spans = all_spans(document)
    assert [span.kind for span in spans] == [SpanKind.PARAGRAPH, SpanKind.TABLE]
    rows = [[cell.strip() for cell in line.split("|")] for line in spans[1].text.split("\n")]
    assert [[cell for cell in row if cell] for row in rows] == [["A", "B"], ["C", "D"]]


# --- DOCX details ------------------------------------------------------------------------------


def test_docx_tables_keep_their_place_and_sections_start_pages(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=2, cols=2)
    for index, text in enumerate(["A", "B", "C", "D"]):
        table.cell(index // 2, index % 2).text = text
    document.add_paragraph("After the table.")
    document.add_paragraph("")
    document.add_section()
    document.add_heading("Next section", level=1)
    document.add_paragraph("A bullet", style="List Bullet")
    path = tmp_path / "tables.docx"
    document.save(str(path))

    result = ingest(path)
    assert [page.number for page in result.pages] == [1, 2]
    assert kinds_and_text(result.pages[0].spans) == [
        ("paragraph", "Before the table."),
        ("table", "A | B\nC | D"),
        ("paragraph", "After the table."),
    ]
    assert kinds_and_text(result.pages[1].spans) == [
        ("heading", "Next section"),
        ("list_item", "A bullet"),
    ]


def test_docx_starts_a_page_every_forty_paragraphs(tmp_path: Path) -> None:
    document = Document()
    for index in range(85):
        document.add_paragraph(f"Line {index + 1}")
    path = tmp_path / "long.docx"
    document.save(str(path))

    result = ingest(path)
    assert [len(page.spans) for page in result.pages] == [40, 40, 5]
    assert result.pages[1].spans[0].span_id == "p2.s1"
    assert result.pages[1].spans[0].text == "Line 41"


def test_docx_trailing_page_break_adds_no_empty_page(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Only page.")
    document.add_page_break()
    path = tmp_path / "trailing.docx"
    document.save(str(path))

    assert [page.number for page in ingest(path).pages] == [1]


def test_docx_jpeg_keeps_its_extension(tmp_path: Path) -> None:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 48), False)
    pixmap.clear_with(180)
    picture = tmp_path / "photo.jpg"
    pixmap.save(str(picture))
    document = Document()
    document.add_paragraph("Caption")
    document.add_picture(str(picture))
    path = tmp_path / "photo.docx"
    document.save(str(path))

    result = ingest(path, image_dir=tmp_path / "out")
    image = result.pages[0].spans[1]
    assert image.kind is SpanKind.IMAGE
    assert image.image_path == f"{result.doc_id}/p1_img1.jpeg"
    assert (tmp_path / "out" / image.image_path).read_bytes().startswith(b"\xff\xd8\xff")


# --- PPTX details ------------------------------------------------------------------------------


def test_pptx_reading_order_tables_groups_and_levels(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    shapes = slide.shapes
    low = shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    low.text_frame.text = "Bottom text"
    top = shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    top.text_frame.text = "Top text"
    top.text_frame.add_paragraph().text = ""
    nested = top.text_frame.add_paragraph()
    nested.text = "Indented point"
    nested.level = 1
    grid = shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for index, text in enumerate(["A", "B", "C", "D"]):
        grid.cell(index // 2, index % 2).text = text
    group = shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(4), Inches(4), Inches(1)
    ).text_frame.text = "Grouped text"
    path = tmp_path / "shapes.pptx"
    presentation.save(str(path))

    result = ingest(path)
    assert kinds_and_text(result.pages[0].spans) == [
        ("paragraph", "Top text"),
        ("list_item", "Indented point"),
        ("table", "A | B\nC | D"),
        ("paragraph", "Grouped text"),
        ("paragraph", "Bottom text"),
    ]
    assert result.pages[0].spans[0].bbox == (72.0, 72.0, 360.0, 144.0)
