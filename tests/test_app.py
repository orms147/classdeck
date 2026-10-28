"""The local teacher app (``classdeck.app``): routes, jobs, proposals, uploads, static page.

Everything runs offline: the draft uses the template pipeline (``endpoint=offline``) and edits
use a scripted ``MockProvider`` injected through ``create_app(provider_factory=...)``.
"""

from __future__ import annotations

import json
import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from classdeck.app import server
from classdeck.app.server import create_app
from classdeck.config import Settings
from classdeck.lesson import Lesson, LessonStore
from classdeck.llm import MockProvider
from classdeck.llm.base import Message

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "eval" / "fixtures" / "public"
SAMPLE = "gs3-unit1"
UNIT = "ta3.u1"


def make_settings(tmp_path: Path, **extra: Any) -> Settings:
    return Settings(
        workdir=tmp_path / "out",
        curriculum_dir=ROOT / "curriculum",
        provider="mock",
        cache_dir=None,
        **extra,
    )


def block_by_kind(lesson: Lesson, kind: str, *, min_items: int = 0) -> Any:
    for block in lesson.ordered_blocks():
        if block.kind == kind and len(getattr(block.content, "items", [])) >= min_items:
            return block
    raise AssertionError(f"no {kind} block with {min_items} items")


def drop_last_word_reply(store: LessonStore, lesson_id: str) -> Callable[[list[Message]], str]:
    """A scripted model: reply with the first vocabulary block minus its last item."""

    def handler(_: list[Message]) -> str:
        lesson = store.load(lesson_id)
        block = block_by_kind(lesson, "vocabulary", min_items=2).model_dump(mode="json")
        block["content"]["items"] = block["content"]["items"][:-1]
        return json.dumps(
            {
                "replacements": {block["id"]: block},
                "added_blocks": [],
                "removed_block_ids": [],
                "rationale": "Dropped the last word.",
            }
        )

    return handler


class Harness:
    """A client plus the pieces tests poke at."""

    def __init__(self, tmp_path: Path, **kwargs: Any) -> None:
        self.settings = make_settings(tmp_path)
        self.store = LessonStore(self.settings.workdir)
        self.provider: MockProvider | None = None
        self.client = TestClient(
            create_app(
                self.settings,
                provider_factory=lambda _settings, endpoint: (
                    None if endpoint == "offline" else self.provider
                ),
                sample_dir=SAMPLE_DIR,
                **kwargs,
            )
        )

    def create_lesson(self, *, pictures: bool = False, **form: Any) -> dict[str, Any]:
        data = {
            "sample": SAMPLE,
            "grade": 3,
            "unit": UNIT,
            "lesson": 1,
            "period": 1,
            "endpoint": "offline",
            "pictures": str(pictures).lower(),
            **form,
        }
        response = self.client.post("/api/jobs", data=data)
        assert response.status_code == 200, response.text
        return self.wait(response.json()["job_id"])

    def wait(self, job_id: str) -> dict[str, Any]:
        deadline = time.time() + 120
        while time.time() < deadline:
            job = self.client.get(f"/api/jobs/{job_id}").json()
            if job["status"] != "running":
                return job
            time.sleep(0.1)
        raise AssertionError("the job did not finish")


@pytest.fixture
def app_env(tmp_path: Path) -> Harness:
    return Harness(tmp_path)


@pytest.fixture
def lesson_id(app_env: Harness) -> str:
    job = app_env.create_lesson()
    assert job["status"] == "done", job
    return job["lesson_id"]


# --- config ------------------------------------------------------------------------------------


def test_config_lists_endpoints_units_themes_and_samples(tmp_path: Path) -> None:
    settings = make_settings(tmp_path, deepseek_api_key="sk-secret-value-1234")
    client = TestClient(create_app(settings, sample_dir=SAMPLE_DIR))
    response = client.get("/api/config")
    assert response.status_code == 200
    config = response.json()

    endpoints = {e["name"]: e for e in config["endpoints"]}
    assert endpoints["offline"]["configured"] is True
    assert endpoints["deepseek"]["configured"] is True
    assert endpoints["meta"]["configured"] is False
    assert endpoints["openrouter"]["configured"] is False
    assert "sk-secret-value-1234" not in response.text

    grades = {g["grade"]: g for g in config["grades"]}
    assert set(grades) == set(range(1, 10))
    assert grades[1]["units"] == [] and grades[7]["units"] == []
    assert grades[2]["lessons"] == 3 and grades[2]["periods"] == 2
    assert grades[8]["unit_shape"] == "secondary" and grades[8]["lessons"] == 7
    assert grades[9]["units"][0]["lessons"] == 7
    assert grades[3]["units"][0] == {
        "id": UNIT,
        "title": grades[3]["units"][0]["title"],
        "lessons": 3,
    }
    assert grades[3]["units"][0]["title"]
    assert {t["id"] for t in config["themes"]} >= {"kids", "plain"}
    assert all(t["label"] for t in config["themes"])
    assert any(s["id"] == SAMPLE and s["unit_id"] == UNIT for s in config["samples"])
    assert config["default_endpoint"] == "offline"
    assert any(p["id"] == "reduce_words" and p["takes_number"] for p in config["presets"])
    assert config["limits"]["upload_mb"] == 30


def test_unconfigured_endpoint_is_refused(app_env: Harness) -> None:
    response = app_env.client.post(
        "/api/jobs",
        data={"sample": SAMPLE, "grade": 3, "unit": UNIT, "endpoint": "meta"},
    )
    assert response.status_code == 400
    assert "API key" in response.json()["error"]


# --- jobs, lessons, exports ---------------------------------------------------------------------


def test_job_runs_the_pipeline_and_saves_the_lesson(app_env: Harness) -> None:
    job = app_env.create_lesson(pictures=True)
    assert job["status"] == "done", job
    assert job["stages_done"] == ["ingest", "outline", "draft", "pictures", "export"]
    assert job["error"] is None

    listing = app_env.client.get("/api/lessons").json()["lessons"]
    assert [row["id"] for row in listing] == [job["lesson_id"]]
    assert listing[0]["grade"] == 3 and listing[0]["unit_id"] == UNIT

    exports = app_env.settings.workdir / "exports" / job["lesson_id"]
    assert {p.name for p in exports.iterdir()} == {"lesson.html"}


def test_job_without_pictures_skips_that_stage(app_env: Harness) -> None:
    job = app_env.create_lesson(pictures=False)
    assert job["stages_skipped"] == ["pictures"]


def test_unknown_unit_and_bad_form_values_are_400(app_env: Harness) -> None:
    base = {"sample": SAMPLE, "grade": 3, "unit": UNIT, "endpoint": "offline"}
    bad = ({"unit": "ta3.u99"}, {"unit": "x"}, {"unit": "0"}, {"grade": 10}, {"lesson": 9})
    for change in (*bad, {"period": 3}, {"theme": "x"}, {"model": "two words"}):
        response = app_env.client.post("/api/jobs", data={**base, **change})
        assert response.status_code == 400, change
        assert "error" in response.json()
    assert app_env.client.post("/api/jobs", data={**base, "sample": "nope"}).status_code == 400
    assert app_env.client.get("/api/jobs/missing").status_code == 404


def test_lesson_endpoint_returns_blocks_for_the_panes(app_env: Harness, lesson_id: str) -> None:
    body = app_env.client.get(f"/api/lessons/{lesson_id}").json()
    lesson, summary = body["lesson"], body["summary"]
    assert lesson["id"] == lesson_id
    rows = summary["blocks"]
    assert rows and {row["id"] for row in rows} == set(lesson["blocks"])
    first = rows[0]
    assert {"kind", "phase", "label", "provenance", "status", "citations"} <= set(first)
    assert summary["can_undo"] is False and summary["can_redo"] is False
    assert any(row["citations"] for row in rows), "some block quotes the source with a page number"
    cited = next(c for row in rows for c in row["citations"])
    assert cited["page"] >= 1 and cited["quote"]
    assert all(r["phase"] in {s["phase"] for s in summary["sections"]} for r in rows)


def test_deck_endpoint_renders_html_with_block_ids(app_env: Harness, lesson_id: str) -> None:
    response = app_env.client.get(f"/api/lessons/{lesson_id}/deck", params={"theme": "plain"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "data-block-id" in response.text
    assert app_env.client.get(f"/api/lessons/{lesson_id}/deck?theme=nope").status_code == 400


@pytest.mark.parametrize(
    ("fmt", "content_type"),
    [
        ("html", "text/html"),
        ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ],
)
def test_export_endpoints_return_the_right_content_type(
    app_env: Harness, lesson_id: str, fmt: str, content_type: str
) -> None:
    response = app_env.client.get(f"/api/lessons/{lesson_id}/export/{fmt}")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(content_type)
    assert f"Grade-3-Lesson-1-Period-1.{fmt}" in response.headers["content-disposition"]
    assert lesson_id not in response.headers["content-disposition"]
    assert len(response.content) > 1000
    if fmt in ("pptx", "docx"):
        assert response.content[:2] == b"PK"


def test_unknown_lesson_and_format_are_json_404(app_env: Harness, lesson_id: str) -> None:
    for url in (
        "/api/lessons/l_nope",
        "/api/lessons/..%2Fx",
        f"/api/lessons/{lesson_id}/export/zip",
    ):
        response = app_env.client.get(url)
        assert response.status_code == 404
        assert "error" in response.json()


def test_check_and_illustrate(app_env: Harness, lesson_id: str) -> None:
    checked = app_env.client.post(f"/api/lessons/{lesson_id}/check", json={})
    assert checked.status_code == 200
    assert checked.json()["lesson"]["id"] == lesson_id

    plan = app_env.client.post(
        f"/api/lessons/{lesson_id}/illustrate",
        json={"max": 2, "dry_run": True, "endpoint": "offline"},
    ).json()
    assert plan["report"]["planned"] == 2
    assert not (app_env.settings.workdir / "images" / "generated").exists()

    done = app_env.client.post(
        f"/api/lessons/{lesson_id}/illustrate", json={"max": 2, "endpoint": "offline"}
    ).json()
    assert done["report"]["generated"] == 2
    pictures = [
        item
        for block in done["lesson"]["blocks"].values()
        if block["kind"] == "vocabulary"
        for item in block["content"]["items"]
        if item.get("generated_picture")
    ]
    assert len(pictures) >= 2


# --- proposals, accept, undo, redo --------------------------------------------------------------


def propose(env: Harness, lesson_id: str, **extra: Any) -> Any:
    lesson = env.store.load(lesson_id)
    block = block_by_kind(lesson, "vocabulary", min_items=2)
    env.provider = MockProvider(handler=drop_last_word_reply(env.store, lesson_id))
    return env.client.post(
        f"/api/lessons/{lesson_id}/proposals",
        json={"block_ids": [block.id], "instruction": "Drop the last word", **extra},
    )


def test_proposal_accept_undo_redo(app_env: Harness, lesson_id: str) -> None:
    before = app_env.store.load(lesson_id)
    block = block_by_kind(before, "vocabulary", min_items=2)
    words = [item.word for item in block.content.items]

    response = propose(app_env, lesson_id)
    assert response.status_code == 200, response.text
    proposal = response.json()
    change = proposal["changes"][0]
    assert change["op"] == "replace" and change["block_id"] == block.id
    assert [i["word"] for i in change["before"]["content"]["items"]] == words
    assert [i["word"] for i in change["after"]["content"]["items"]] == words[:-1]
    assert proposal["rationale"] == "Dropped the last word."
    assert isinstance(proposal["affected"], list) and isinstance(proposal["checks_preview"], list)
    # Nothing is saved until the teacher accepts.
    assert app_env.store.load(lesson_id).model_dump() == before.model_dump()

    accepted = app_env.client.post(f"/api/lessons/{lesson_id}/proposals/{proposal['id']}/accept")
    assert accepted.status_code == 200, accepted.text
    body = accepted.json()
    assert body["summary"]["can_undo"] is True and body["summary"]["can_redo"] is False
    new_words = [i["word"] for i in body["lesson"]["blocks"][block.id]["content"]["items"]]
    assert new_words == words[:-1]
    assert len(body["lesson"]["revisions"]) == 1
    assert [i.word for i in app_env.store.load(lesson_id).blocks[block.id].content.items] == words[
        :-1
    ]
    # The proposal is spent.
    again = app_env.client.post(f"/api/lessons/{lesson_id}/proposals/{proposal['id']}/accept")
    assert again.status_code == 404

    undone = app_env.client.post(f"/api/lessons/{lesson_id}/undo").json()
    assert undone["summary"]["can_undo"] is False and undone["summary"]["can_redo"] is True
    assert [i["word"] for i in undone["lesson"]["blocks"][block.id]["content"]["items"]] == words

    redone = app_env.client.post(f"/api/lessons/{lesson_id}/redo").json()
    assert redone["summary"]["can_undo"] is True and redone["summary"]["can_redo"] is False
    assert [i["word"] for i in redone["lesson"]["blocks"][block.id]["content"]["items"]] == words[
        :-1
    ]

    assert app_env.client.post(f"/api/lessons/{lesson_id}/redo").status_code == 409


def test_undo_with_no_edits_is_409(app_env: Harness, lesson_id: str) -> None:
    response = app_env.client.post(f"/api/lessons/{lesson_id}/undo")
    assert response.status_code == 409
    assert "error" in response.json()


def test_discarded_proposal_cannot_be_accepted(app_env: Harness, lesson_id: str) -> None:
    proposal = propose(app_env, lesson_id).json()
    url = f"/api/lessons/{lesson_id}/proposals/{proposal['id']}"
    assert app_env.client.delete(url).json() == {"ok": True}
    assert app_env.client.post(url + "/accept").status_code == 404


def test_proposal_validation(app_env: Harness, lesson_id: str) -> None:
    url = f"/api/lessons/{lesson_id}/proposals"
    block_id = next(iter(app_env.store.load(lesson_id).blocks))
    # No model behind the default endpoint.
    offline = app_env.client.post(
        url, json={"block_ids": [block_id], "instruction": "x", "endpoint": "offline"}
    )
    assert offline.status_code == 400 and "model" in offline.json()["error"]
    app_env.provider = MockProvider()
    assert (
        app_env.client.post(url, json={"block_ids": ["b_nope"], "instruction": "x"}).status_code
        == 400
    )
    assert (
        app_env.client.post(url, json={"block_ids": [block_id], "instruction": "  "}).status_code
        == 400
    )
    assert app_env.client.post(url, json={"instruction": "x"}).status_code == 422


def test_model_failure_is_502_and_saves_nothing(app_env: Harness, lesson_id: str) -> None:
    before = app_env.store.load(lesson_id).model_dump()
    block = block_by_kind(app_env.store.load(lesson_id), "vocabulary")
    app_env.provider = MockProvider(responses=[])  # raises LLMError on the first call
    response = app_env.client.post(
        f"/api/lessons/{lesson_id}/proposals",
        json={"block_ids": [block.id], "instruction": "Make it shorter"},
    )
    assert response.status_code == 502
    assert "error" in response.json()
    assert app_env.store.load(lesson_id).model_dump() == before


def test_unusable_model_reply_is_422(app_env: Harness, lesson_id: str) -> None:
    block = block_by_kind(app_env.store.load(lesson_id), "vocabulary")
    empty = json.dumps(
        {"replacements": {}, "added_blocks": [], "removed_block_ids": [], "rationale": "no"}
    )
    app_env.provider = MockProvider(handler=lambda _: empty)
    response = app_env.client.post(
        f"/api/lessons/{lesson_id}/proposals",
        json={"block_ids": [block.id], "instruction": "Do something"},
    )
    assert response.status_code in (422, 502), response.text
    assert "error" in response.json()


# --- uploads ------------------------------------------------------------------------------------


def upload(env: Harness, name: str, content: bytes) -> Any:
    return env.client.post(
        "/api/jobs",
        data={"grade": 3, "unit": UNIT, "endpoint": "offline"},
        files={"file": (name, content, "application/octet-stream")},
    )


def test_upload_rejects_other_formats(app_env: Harness) -> None:
    response = upload(app_env, "notes.txt", b"hello")
    assert response.status_code == 400
    assert "PDF" in response.json()["error"]
    assert not (app_env.settings.workdir / "uploads").exists() or not list(
        (app_env.settings.workdir / "uploads").iterdir()
    )


def test_upload_rejects_a_file_that_is_not_what_its_name_says(app_env: Harness) -> None:
    assert upload(app_env, "fake.pdf", b"not a pdf").status_code == 400
    assert upload(app_env, "empty.pdf", b"").status_code == 400


def test_upload_rejects_oversize_files(app_env: Harness, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(server, "MAX_UPLOAD_BYTES", 2048)
    response = upload(app_env, "big.pdf", b"%PDF-1.4\n" + b"0" * 5000)
    assert response.status_code == 413
    assert "30 MB" in response.json()["error"]
    uploads = app_env.settings.workdir / "uploads"
    assert not uploads.exists() or not list(uploads.iterdir())


def test_upload_is_saved_under_a_generated_name_and_ingested(app_env: Harness) -> None:
    content = (SAMPLE_DIR / f"{SAMPLE}.pdf").read_bytes()
    response = upload(app_env, "../../evil name.PDF", content)
    assert response.status_code == 200, response.text
    job = app_env.wait(response.json()["job_id"])
    assert job["status"] == "done", job
    saved = list((app_env.settings.workdir / "uploads").iterdir())
    assert len(saved) == 1
    assert saved[0].suffix == ".pdf" and "evil" not in saved[0].name


# --- the page -----------------------------------------------------------------------------------


def test_static_page_is_served(app_env: Harness) -> None:
    index = app_env.client.get("/")
    assert index.status_code == 200
    assert index.headers["content-type"].startswith("text/html")
    assert "ClassDeck" in index.text and "/static/app.js" in index.text
    for name, kind in (("app.js", "javascript"), ("app.css", "css"), ("i18n.js", "javascript")):
        response = app_env.client.get(f"/static/{name}")
        assert response.status_code == 200 and kind in response.headers["content-type"]
    assert app_env.client.get("/static/../server.py").status_code in (400, 404)


def test_only_loopback_hosts_are_served(app_env: Harness) -> None:
    assert app_env.client.get("/api/config", headers={"host": "evil.example"}).status_code == 403
    assert app_env.client.get("/api/config", headers={"host": "127.0.0.1:8765"}).status_code == 200
    assert app_env.client.get("/api/config", headers={"host": "localhost:9"}).status_code == 200


def test_ui_strings_cover_the_keys_the_page_uses() -> None:
    """Every literal ``t('key')`` in app.js exists in both languages."""
    import re

    static = Path(server.__file__).parent / "static"
    script = (static / "app.js").read_text(encoding="utf-8")
    strings = (static / "i18n.js").read_text(encoding="utf-8")
    keys = set(re.findall(r"\bt\('([a-z0-9_.]+)'", script)) - {
        "stage.",
        "prov.",
        "status.",
        "dl.",
        "prop.op.",
    }
    vi, en = strings.split("\n  en: {")
    for key in sorted(keys):
        needle = f"'{key}':" if "." in key else f"{key}:"
        assert needle in vi, f"vi is missing {key}"
        assert needle in en, f"en is missing {key}"


def test_serve_refuses_non_loopback_hosts() -> None:
    with pytest.raises(ValueError, match="loopback"):
        server.serve(host="0.0.0.0", open_browser=False)


# --- grade 9 (secondary) -------------------------------------------------------------------------


def test_config_lists_grade_9_as_a_secondary_grade_with_seven_lessons(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    config = TestClient(create_app(settings, sample_dir=SAMPLE_DIR)).get("/api/config").json()
    grades = {g["grade"]: g for g in config["grades"]}
    assert grades[9]["unit_shape"] == "secondary" and grades[3]["unit_shape"] == "primary"
    assert [u["id"] for u in grades[9]["units"]][:2] == ["ta9.u1", "ta9.u2"]
    assert all(u["lessons"] == 7 for u in grades[9]["units"])


def test_a_grade_9_job_takes_lessons_up_to_7(app_env: Harness) -> None:
    base = {"sample": SAMPLE, "grade": 9, "unit": "ta9.u1", "endpoint": "offline"}
    response = app_env.client.post("/api/jobs", data={**base, "lesson": 8})
    assert response.status_code == 400 and "1-7" in response.json()["error"]
    job_id = app_env.client.post("/api/jobs", data={**base, "lesson": 7}).json()["job_id"]
    for _ in range(100):
        job = app_env.client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            break
        time.sleep(0.05)
    # the sample is a primary text PDF: no section banners, so the job reports it cleanly
    assert job["status"] == "error" and "LOOKING BACK & PROJECT banner" in job["error"]


# --- editing on the slide (docs/output-formats.md) -----------------------------------------------


def first_vocab_field(app_env: Harness, lesson_id: str, name: str) -> tuple[str, str, str]:
    """``(block id, path, current text)`` of the first vocabulary item's ``name``."""
    lesson = app_env.store.load(lesson_id)
    block = next(b for b in lesson.ordered_blocks() if b.kind == "vocabulary")
    item = block.content.items[0]  # type: ignore[union-attr]
    return block.id, f"content/items/0/{name}", getattr(item, name)


def test_the_deck_is_served_with_its_lesson_id_and_the_edit_controls(
    app_env: Harness, lesson_id: str
) -> None:
    html = app_env.client.get(f"/api/lessons/{lesson_id}/deck").text
    assert f'<meta name="classdeck-lesson" content="{lesson_id}">' in html
    assert 'id="edit-toggle"' in html and 'id="edit-save"' in html
    assert 'data-field="content/items/0/word"' in html


def test_saving_edits_is_one_revision_that_undo_takes_back(
    app_env: Harness, lesson_id: str
) -> None:
    block_id, word_path, word = first_vocab_field(app_env, lesson_id, "word")
    _, example_path, _ = first_vocab_field(app_env, lesson_id, "example")
    lesson = app_env.store.load(lesson_id)
    activity = next(b for b in lesson.ordered_blocks() if b.kind == "activity")
    response = app_env.client.post(
        f"/api/lessons/{lesson_id}/edits",
        json={
            "edits": [
                {"block_id": block_id, "field": word_path, "value": word + "s"},
                {"block_id": block_id, "field": example_path, "value": "A new   example."},
                {"block_id": activity.id, "field": "content/title", "value": "A new title"},
            ],
            "theme": "plain",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert set(body["changed"]) == {block_id, activity.id}
    assert body["exported"] is True
    assert isinstance(body["notices"], list)
    assert body["summary"]["can_undo"] is True
    assert len(body["lesson"]["revisions"]) == 1
    assert body["lesson"]["revisions"][0]["actor"] == "teacher"

    saved = app_env.store.load(lesson_id)
    assert len(saved.revisions) == 1
    assert saved.blocks[block_id].content.items[0].word == word + "s"
    assert saved.blocks[block_id].content.items[0].example == "A new example."
    assert saved.blocks[activity.id].content.title == "A new title"
    deck = app_env.settings.workdir / "exports" / lesson_id / "lesson.html"
    assert "A new title" in deck.read_text(encoding="utf-8")
    assert 'data-theme="plain"' in deck.read_text(encoding="utf-8")

    undone = app_env.client.post(f"/api/lessons/{lesson_id}/undo").json()
    assert undone["lesson"]["blocks"][block_id]["content"]["items"][0]["word"] == word
    assert undone["lesson"]["blocks"][activity.id]["content"]["title"] == activity.content.title
    assert undone["summary"]["can_undo"] is False and undone["summary"]["can_redo"] is True


def test_a_bare_list_of_edits_is_accepted_and_uses_the_default_theme(
    app_env: Harness, lesson_id: str
) -> None:
    block_id, path, meaning = first_vocab_field(app_env, lesson_id, "meaning_vi")
    response = app_env.client.post(
        f"/api/lessons/{lesson_id}/edits",
        json=[{"block_id": block_id, "field": path, "value": (meaning or "") + " nhe"}],
    )
    assert response.status_code == 200, response.text
    deck = app_env.settings.workdir / "exports" / lesson_id / "lesson.html"
    assert 'data-theme="sunny"' in deck.read_text(encoding="utf-8")


def test_edits_that_change_nothing_make_no_revision(app_env: Harness, lesson_id: str) -> None:
    block_id, path, word = first_vocab_field(app_env, lesson_id, "word")
    body = app_env.client.post(
        f"/api/lessons/{lesson_id}/edits",
        json={"edits": [{"block_id": block_id, "field": path, "value": word}]},
    ).json()
    assert body["changed"] == [] and body["exported"] is False
    assert body["summary"]["can_undo"] is False
    assert (
        app_env.client.post(f"/api/lessons/{lesson_id}/edits", json={"edits": []}).status_code
        == 200
    )


def test_a_refused_edit_is_a_422_that_names_the_field_and_changes_nothing(
    app_env: Harness, lesson_id: str
) -> None:
    block_id, path, word = first_vocab_field(app_env, lesson_id, "word")
    before = app_env.store.load(lesson_id)
    bad = [
        ({"block_id": block_id, "field": path, "value": ""}, "cannot be empty", path),
        ({"block_id": block_id, "field": "content/items/0/nope", "value": "x"}, "unknown", None),
        ({"block_id": block_id, "field": "provenance", "value": "generated"}, "not editable", None),
        ({"block_id": "b_nope", "field": path, "value": "x"}, "no block", path),
    ]
    for edit, message, field in bad:
        response = app_env.client.post(
            f"/api/lessons/{lesson_id}/edits",
            json={"edits": [{"block_id": block_id, "field": path, "value": word + "!"}, edit]},
        )
        assert response.status_code == 422, response.text
        body = response.json()
        assert message in body["error"]
        assert body["block_id"] == edit["block_id"] and body["field"] == (field or edit["field"])
    assert app_env.store.load(lesson_id).model_dump() == before.model_dump()
    response = app_env.client.post(
        f"/api/lessons/{lesson_id}/edits", json={"edits": [{"block_id": block_id}]}
    )
    assert response.status_code == 422
    unknown_theme = app_env.client.post(
        f"/api/lessons/{lesson_id}/edits", json={"edits": [], "theme": "nope"}
    )
    assert unknown_theme.status_code == 400
    assert app_env.client.post("/api/lessons/l_nope/edits", json={"edits": []}).status_code == 404


# --- teacher-first simplification (docs/local-app.md) ---------------------------------------------

SECRET = "sk-or-v1-0123456789abcdefSECRETTAIL"


def secondary_docx(path: Path) -> bytes:
    """A made-up grade 6-8 unit page with the section banners the segmentation looks for."""
    from docx import Document

    document = Document()
    document.add_heading("UNIT 3 MY NEIGHBOURHOOD", level=1)
    document.add_heading("THIS UNIT INCLUDES:", level=2)
    document.add_paragraph("Vocabulary | places in a neighbourhood")
    document.add_paragraph("Grammar | comparative adjectives")
    document.add_heading("GETTING STARTED", level=1)
    document.add_paragraph("1. Listen and read.")
    document.add_paragraph("Lan: Where is the library?")
    document.add_paragraph("Nam: It is next to the post office.")
    document.add_paragraph("2. Fill in each blank with ONE word.")
    document.add_paragraph("1. The library is next to the ___ office.")
    document.add_heading("A CLOSER LOOK 1", level=1)
    document.add_paragraph("1. Match the words with their meanings.")
    document.save(str(path))
    return path.read_bytes()


def body_of(html: str, class_name: str) -> str:
    import re

    match = re.search(rf'class="{class_name}[^"]*"[^>]*>(.*?)</', html, re.S)
    assert match, class_name
    return match.group(1)


def test_settings_start_empty_and_never_return_the_key(app_env: Harness) -> None:
    shown = app_env.client.get("/api/settings")
    assert shown.status_code == 200
    body = shown.json()
    assert body["key_set"] is False and body["key_hint"] is None
    assert body["model"] == "openai/gpt-6-luna" and body["default_theme"] == "sunny"
    assert body["workdir"] == str(app_env.settings.workdir.resolve())

    saved = app_env.client.put(
        "/api/settings",
        json={"api_key": f"  {SECRET} ", "model": "vendor/newer-model", "default_theme": "kids"},
    )
    assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body["key_set"] is True and body["key_hint"] == "sk-…TAIL"
    assert body["model"] == "vendor/newer-model" and body["default_theme"] == "kids"
    for path in ("/api/settings", "/api/config", "/api/lessons"):
        assert SECRET not in app_env.client.get(path).text
    assert SECRET not in saved.text and "0123456789" not in saved.text
    config = app_env.client.get("/api/config").json()
    assert config["key_set"] is True and config["default_theme"] == "kids"


def test_settings_are_stored_in_the_user_settings_file(app_env: Harness) -> None:
    from classdeck.config import load_user_settings, user_settings_path

    app_env.client.put("/api/settings", json={"api_key": SECRET, "model": "vendor/m"})
    assert load_user_settings() == {"openrouter_api_key": SECRET, "model": "vendor/m"}
    assert user_settings_path().is_file()
    # a value left out stays; an empty string removes it and brings the default back
    app_env.client.put("/api/settings", json={"default_theme": "kids"})
    assert load_user_settings()["openrouter_api_key"] == SECRET
    body = app_env.client.put("/api/settings", json={"api_key": "", "model": ""}).json()
    assert body["key_set"] is False and body["model"] == "openai/gpt-6-luna"
    assert load_user_settings() == {"default_theme": "kids"}


def test_settings_reject_bad_values(app_env: Harness) -> None:
    for body in (
        {"api_key": "two words"},
        {"api_key": "k" * 400},
        {"model": "has space"},
        {"default_theme": "nope"},
    ):
        response = app_env.client.put("/api/settings", json=body)
        assert response.status_code == 400, body
        assert "error" in response.json()
    assert app_env.client.get("/api/settings").json()["key_set"] is False


def test_the_saved_key_makes_the_default_model_available(app_env: Harness) -> None:
    assert app_env.client.get("/api/config").json()["default_endpoint"] == "offline"
    app_env.client.put("/api/settings", json={"api_key": SECRET})
    config = app_env.client.get("/api/config").json()
    endpoints = {e["name"]: e for e in config["endpoints"]}
    assert endpoints["openrouter"]["configured"] is True


def test_settings_test_sends_one_request_and_shows_the_reply(app_env: Harness) -> None:
    app_env.provider = MockProvider(["pong"], model="vendor/newer-model")
    result = app_env.client.post(
        "/api/settings/test", json={"api_key": SECRET, "model": "vendor/newer-model"}
    )
    assert result.status_code == 200, result.text
    assert result.json() == {"ok": True, "model": "vendor/newer-model", "reply": "pong"}
    assert len(app_env.provider.calls) == 1
    assert SECRET not in result.text
    # nothing is saved by testing
    assert app_env.client.get("/api/settings").json()["key_set"] is False


def test_settings_test_reports_errors_without_the_key(app_env: Harness) -> None:
    from classdeck.llm.base import LLMError

    def refuse(_: list[Message]) -> str:
        raise LLMError(f"401: Incorrect API key provided: {SECRET}")

    app_env.provider = MockProvider(handler=refuse)
    result = app_env.client.post("/api/settings/test", json={"api_key": SECRET})
    assert result.status_code == 200
    body = result.json()
    assert body["ok"] is False and "401" in body["error"]
    assert SECRET not in result.text


def test_settings_test_without_a_key_asks_for_one(tmp_path: Path) -> None:
    settings = make_settings(tmp_path).model_copy(update={"provider": "api"})
    client = TestClient(create_app(settings, sample_dir=SAMPLE_DIR))
    body = client.post("/api/settings/test", json={}).json()
    assert body["ok"] is False and "key" in body["error"].lower()


def test_open_folder_opens_the_work_folder_for_local_requests(tmp_path: Path) -> None:
    opened: list[Path] = []
    settings = make_settings(tmp_path)
    client = TestClient(create_app(settings, sample_dir=SAMPLE_DIR, folder_opener=opened.append))
    response = client.post("/api/open-folder")
    assert response.status_code == 200
    assert opened == [settings.workdir.resolve()]
    assert response.json()["path"] == str(settings.workdir.resolve())
    remote = TestClient(
        create_app(settings, sample_dir=SAMPLE_DIR, folder_opener=opened.append),
        client=("203.0.113.9", 5000),
    )
    assert remote.post("/api/open-folder").status_code == 403
    assert len(opened) == 1


def test_a_grade_2_job_takes_a_unit_number_and_builds_the_unit_from_the_pages(
    app_env: Harness,
) -> None:
    job = app_env.create_lesson(grade=2, unit="1")
    assert job["status"] == "done", job
    row = app_env.client.get("/api/lessons").json()["lessons"][0]
    assert row["grade"] == 2 and row["unit_id"] == "ta2.u1"
    assert row["title"] == "Grade 2 · Lesson 1 · Period 1"
    assert row["unit"] == "Unit 1 · Hello"
    payload = app_env.client.get(f"/api/lessons/{row['id']}").json()
    assert payload["lesson"]["title"] == "Grade 2 · Lesson 1 · Period 1"
    assert payload["summary"]["unit_line"] == "Unit 1 · Hello"
    # the later calls rebuild the same unit from the saved pages
    assert app_env.client.post(f"/api/lessons/{row['id']}/check", json={}).status_code == 200
    assert app_env.client.get(f"/api/lessons/{row['id']}/export/docx").status_code == 200


def test_the_old_unit_id_is_still_accepted_for_a_unit_in_the_curriculum(app_env: Harness) -> None:
    job = app_env.create_lesson(unit=UNIT)
    assert job["status"] == "done", job
    assert app_env.client.get("/api/lessons").json()["lessons"][0]["unit_id"] == UNIT
    assert app_env.client.post("/api/jobs", data={"grade": 4, "unit": UNIT}).status_code == 400


def test_a_unit_that_is_not_in_the_curriculum_file_is_built_from_the_pages(
    app_env: Harness,
) -> None:
    job = app_env.create_lesson(unit="14")
    assert job["status"] == "done", job
    row = app_env.client.get("/api/lessons").json()["lessons"][0]
    assert row["unit_id"] == "ta3.u14" and row["unit_no"] == 14
    assert row["unit"].startswith("Unit 14")


def test_the_deck_names_the_lesson_without_internal_ids(app_env: Harness) -> None:
    job = app_env.create_lesson(grade=2, unit="1")
    html = app_env.client.get(f"/api/lessons/{job['lesson_id']}/deck").text
    assert "<title>Grade 2 · Lesson 1 · Period 1 — Unit 1 · Hello</title>" in html
    assert "Unit 1 · Hello" in body_of(html, "cover__title")
    assert body_of(html, "cover__meta") == "Grade 2 · Lesson 1 · Period 1"
    assert "ta2.u1" not in body_of(html, "cover__title") + body_of(html, "cover__meta")
    assert "ta2.u1" not in html.split("</title>")[0]


@pytest.mark.parametrize("grade", [6, 7, 8])
def test_secondary_grades_without_a_curriculum_file_run_end_to_end(
    app_env: Harness, tmp_path: Path, grade: int
) -> None:
    content = secondary_docx(tmp_path / "unit.docx")
    response = app_env.client.post(
        "/api/jobs",
        data={"grade": grade, "unit": "3", "lesson": 1, "endpoint": "offline"},
        files={"file": ("unit.docx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    job = app_env.wait(response.json()["job_id"])
    assert job["status"] == "done", job
    row = app_env.client.get("/api/lessons").json()["lessons"][0]
    assert row["title"] == f"Grade {grade} · Lesson 1 · Getting Started"
    assert row["unit"] == "Unit 3 · My neighbourhood"
    assert row["period"] is None
    html = app_env.client.get(f"/api/lessons/{row['id']}/deck").text
    title = f"Grade {grade} · Lesson 1 · Getting Started — Unit 3 · My neighbourhood"
    assert f"<title>{title}</title>" in html
    assert app_env.client.post(f"/api/lessons/{row['id']}/check", json={}).status_code == 200


def test_a_reference_plan_is_for_secondary_grades_only(app_env: Harness, tmp_path: Path) -> None:
    plan = secondary_docx(tmp_path / "plan.docx")
    data = {"sample": SAMPLE, "grade": 3, "unit": "1", "endpoint": "offline"}
    response = app_env.client.post(
        "/api/jobs", data=data, files={"reference": ("plan.docx", plan, "application/octet-stream")}
    )
    assert response.status_code == 400 and "6-9" in response.json()["error"]
    bad = app_env.client.post(
        "/api/jobs",
        data={**data, "grade": 7},
        files={"reference": ("plan.pdf", b"%PDF-1.4", "application/octet-stream")},
    )
    assert bad.status_code == 400 and "DOCX" in bad.json()["error"]


def test_a_model_override_reaches_the_provider(tmp_path: Path) -> None:
    seen: list[tuple[str, str | None]] = []

    def factory(settings: Settings, endpoint: str | None) -> None:
        seen.append((settings.model, endpoint))

    client = TestClient(
        create_app(make_settings(tmp_path), provider_factory=factory, sample_dir=SAMPLE_DIR)
    )
    response = client.post(
        "/api/jobs",
        data={"sample": SAMPLE, "grade": 3, "unit": "1", "model": "vendor/other-model"},
    )
    assert response.status_code == 200
    for _ in range(100):
        if seen:
            break
        time.sleep(0.05)
    assert seen[0] == ("vendor/other-model", None)


def test_a_scanned_page_that_cannot_be_read_is_a_loud_job_warning(
    app_env: Harness, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pages that fail to read are named in the job's warnings, not dropped silently."""
    from tests.test_ingest_vision import scanned_page_pdf

    monkeypatch.setattr("classdeck.ingest.vision.time.sleep", lambda _s: None)

    app_env.provider = MockProvider(responses=[])  # every call raises LLMError
    pdf = scanned_page_pdf(tmp_path / "scan.pdf", pages=2)
    response = app_env.client.post(
        "/api/jobs",
        data={"grade": 9, "unit": "1", "lesson": 1},
        files={"file": ("scan.pdf", pdf.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    job = app_env.wait(response.json()["job_id"])
    warnings = " ".join(job["warnings"])
    assert "page 1 could not be read: MockProvider has no canned responses" in warnings
    assert "page 2 could not be read" in warnings
    assert "2 scanned page(s) were not read (1, 2)" in warnings
