"""The interactive exercise blocks: schema, text views, HTML deck, PPTX, DOCX and prompt pack."""

from __future__ import annotations

import copy
import json
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pptx import Presentation
from pydantic import ValidationError

from classdeck.english import validators as v
from classdeck.export import export_docx, export_html, export_pptx
from classdeck.grounding import comparable_text
from classdeck.handoff.common import render_lesson_for_prompt
from classdeck.lesson import EXERCISE_KINDS, BlockAdapter, BlockKind, Lesson
from classdeck.lesson.exercises import answer_key, exercise_lines, exercise_title
from tests.exercise_fixtures import (
    CHOICE,
    GAP_FILL,
    LONG_GAP,
    MATCH,
    QUESTIONS,
    READING,
    RULE_BOX,
    TRUE_FALSE,
    make_exercise_lesson,
)
from tests.test_export_html import THEMES_README, Deck, Node, load

ALL = [READING, RULE_BOX, MATCH, GAP_FILL, CHOICE, TRUE_FALSE, LONG_GAP, QUESTIONS]


# --- schema ------------------------------------------------------------------------------------


@pytest.fixture
def lesson() -> Lesson:
    return make_exercise_lesson()


def changed(base: dict[str, Any], **content: Any) -> dict[str, Any]:
    block = copy.deepcopy(base)
    block["content"].update(content)
    return block


def test_the_exercise_kinds_are_block_kinds() -> None:
    assert {kind.value for kind in EXERCISE_KINDS} == {
        "reading",
        "match",
        "gap_fill",
        "choice",
        "true_false",
        "rule_box",
        "question_set",
        "sound_sort",
        "gallery",
        "grid",
        "self_check",
    }
    assert set(BlockKind) >= EXERCISE_KINDS


@pytest.mark.parametrize("block", ALL, ids=[b["id"] for b in ALL])
def test_every_kind_validates_and_round_trips(block: dict[str, Any]) -> None:
    parsed = BlockAdapter.validate_python(block)
    assert parsed.kind == block["kind"]
    again = BlockAdapter.validate_json(parsed.model_dump_json())
    assert again == parsed


def test_lesson_round_trips_through_json(lesson: Lesson) -> None:
    assert Lesson.model_validate_json(lesson.model_dump_json()) == lesson


def test_match_answer_keys_become_integers(lesson: Lesson) -> None:
    answer = lesson.blocks["e_match"].content.answer  # type: ignore[union-attr]
    assert answer == {0: 4, 1: 0, 2: 2, 3: 3, 4: 1}


@pytest.mark.parametrize(
    ("base", "bad"),
    [
        (MATCH, {"answer": {"0": 9, "1": 0, "2": 2, "3": 3, "4": 1}}),  # right index out of range
        (MATCH, {"answer": {"0": 4, "1": 0, "2": 2, "3": 3}}),  # a left item has no answer
        (MATCH, {"answer": {"0": 4, "1": 4, "2": 2, "3": 3, "4": 1}}),  # right item used twice
        (MATCH, {"answer": {"0": 4, "1": 0, "2": 2, "3": 3, "4": 1, "7": 1}}),  # unknown left item
        (MATCH, {"left": ["only one"], "answer": {"0": 0}}),
        (CHOICE, {"items": [{"stem": "Q", "options": ["a", "b"], "answer": "c"}]}),
        (CHOICE, {"items": [{"stem": "Q", "options": ["a", "a"], "answer": "a"}]}),
        (CHOICE, {"items": [{"stem": "Q", "options": ["only"], "answer": "only"}]}),
        (GAP_FILL, {"items": [{"text": "No blank here.", "answer": "x"}]}),
        (GAP_FILL, {"items": [{"text": "One ___ blank.", "answer": ["a", "b"]}]}),
        (GAP_FILL, {"items": [{"text": "A ___ B.", "answer": "missing"}]}),  # not in the word box
        (READING, {"paragraphs": [{"text": "Short text.", "highlight": ["absent"]}]}),
        (READING, {"paragraphs": []}),
        (TRUE_FALSE, {"items": []}),
        (QUESTIONS, {"questions": []}),
    ],
)
def test_bad_answers_are_rejected(base: dict[str, Any], bad: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        BlockAdapter.validate_python(changed(base, **bad))


def test_unknown_fields_are_rejected() -> None:
    with pytest.raises(ValidationError):
        BlockAdapter.validate_python(changed(MATCH, surprise=1))


def test_gap_with_two_blanks_takes_a_list_of_answers() -> None:
    block = BlockAdapter.validate_python(
        changed(
            LONG_GAP,
            items=[{"text": "She ___ to ___ every day.", "answer": ["goes", "school"]}],
        )
    )
    assert block.content.items[0].answers == ["goes", "school"]  # type: ignore[union-attr]


def test_bare_strings_stand_for_paragraphs_and_questions() -> None:
    reading = BlockAdapter.validate_python(changed(READING, paragraphs=["One.", "Two."]))
    assert [p.text for p in reading.content.paragraphs] == ["One.", "Two."]  # type: ignore[union-attr]
    questions = BlockAdapter.validate_python(changed(QUESTIONS, questions=["What?", "Why?"]))
    assert [q.answer for q in questions.content.questions] == ["", ""]  # type: ignore[union-attr]


# --- text views --------------------------------------------------------------------------------


def test_pupil_text_covers_every_new_kind(lesson: Lesson) -> None:
    pupil = {block_id: v.pupil_text_of(block) for block_id, block in lesson.blocks.items()}
    assert "Read the text about city life." in pupil["e_read"]
    assert any("busy city" in text for text in pupil["e_read"])
    assert "Remember!" in pupil["e_rule"]
    assert "She takes the bus every day." in pupil["e_rule"]
    assert {"busy", "calm, with little noise", "very large"} <= set(pupil["e_match"])
    assert "takes" in pupil["e_gap"]  # the word box and the revealed answer
    assert "wakes up" in pupil["e_choice"]
    assert "Mai goes to school by bus." in pupil["e_tf"]
    assert "How do you go to school?" in pupil["e_q"]
    for block_id in pupil:
        assert pupil[block_id], block_id


def test_sample_answers_are_teacher_text_and_glosses_are_left_out(lesson: Lesson) -> None:
    question_set = lesson.blocks["e_q"]
    assert "I go to school by bike." not in v.pupil_text_of(question_set)
    assert "I go to school by bike." in v.english_text_of(question_set)
    glossary_text = " ".join(v.english_text_of(lesson.blocks["e_read"]))
    assert "full of activity" not in glossary_text


def test_comparable_text_is_what_the_page_shows(lesson: Lesson) -> None:
    gap = comparable_text(lesson.blocks["e_gap"])
    assert "Mai ___ the bus to school." in gap
    assert "Complete the sentences" in gap
    assert comparable_text(lesson.blocks["e_match"]).count("full of people") == 1
    assert "I go to school by bike." not in comparable_text(lesson.blocks["e_q"])
    for block in lesson.blocks.values():
        assert comparable_text(block).strip(), block.id


def test_answer_keys(lesson: Lesson) -> None:
    keys = {block_id: answer_key(block) for block_id, block in lesson.blocks.items()}
    assert keys["e_match"][0] == "1-E  busy = full of activity"
    assert keys["e_gap"][0] == "1. Mai takes the bus to school."
    assert keys["e_choice"][0] == "1. B. goes"
    assert keys["e_tf"] == ["1. True", "2. False", "3. True"]
    assert keys["e_q"] == ["1. I go to school by bike.", "2. I live in a village."]
    assert keys["e_read"] == [] and keys["e_rule"] == []


def test_exercise_lines_and_titles(lesson: Lesson) -> None:
    lines = exercise_lines(lesson.blocks["e_match"])
    assert lines[0] == "1. busy" and lines[5] == "A. calm, with little noise"
    assert exercise_title(lesson.blocks["e_match"]) == "Match 2"
    assert exercise_title(lesson.blocks["e_rule"]) == "Remember!"
    assert exercise_title(lesson.blocks["e_read"]) == "Life in a big city"


# --- the HTML deck -----------------------------------------------------------------------------


@pytest.fixture
def deck(tmp_path: Path, lesson: Lesson) -> Deck:
    return load(export_html(lesson, tmp_path / "deck.html", theme="sunny"))


def slides_of(deck: Deck, kind: str) -> list[Node]:
    return deck.root.find("section", "slide", data_kind=kind)


def test_one_slide_per_exercise_and_long_ones_are_split(deck: Deck) -> None:
    types = [(s.attrs["data-slide-type"], s.attrs["data-kind"]) for s in deck.slides()]
    assert types == [
        ("cover", "cover"),
        ("divider", "divider"),
        ("exercise", "reading"),
        ("rule", "rule_box"),
        ("divider", "divider"),
        ("exercise", "match"),
        ("exercise", "gap_fill"),
        ("exercise", "choice"),
        ("exercise", "true_false"),
        ("exercise", "gap_fill"),  # the long one, first part
        ("exercise", "gap_fill"),  # ... and the rest
        ("divider", "divider"),
        ("exercise", "question_set"),
        ("source-map", "source-map"),
    ]
    first, second = slides_of(deck, "gap_fill")[1:]
    assert first.one(class_="ex__part").text() == "(1/2)"
    assert second.one(class_="ex__part").text() == "(2/2)"
    # at most five lines to a slide, cut evenly (8 lines: 4 + 4)
    assert len(first.find(class_="gap-item")) == 4
    assert len(second.find(class_="gap-item")) == 4
    assert [n.text() for n in second.find(class_="gap-item__no")] == ["5", "6", "7", "8"]
    assert not slides_of(deck, "gap_fill")[0].find(class_="ex__part")


def test_exercise_head_shows_the_number_and_the_instruction(deck: Deck) -> None:
    slide = slides_of(deck, "match")[0]
    assert slide.one(class_="ex__no").text() == "2"
    instruction = slide.one(class_="ex__instruction")
    assert instruction.text() == "Match the words with their meanings."
    assert "slide-title" in instruction.classes
    assert slide.one(class_="ex").attrs["data-block-id"] == "e_match"


def test_match_hooks_and_answer_data(deck: Deck) -> None:
    slide = slides_of(deck, "match")[0]
    left, right = slide.one(class_="match__left"), slide.one(class_="match__right")
    assert [n.attrs["data-side"] for n in left.find("button")] == ["left"] * 5
    assert len(right.find("button", "match__item")) == 6  # five answers and a distractor
    assert json.loads(slide.one(class_="match").attrs["data-answer"]) == {
        "0": 4,
        "1": 0,
        "2": 2,
        "3": 3,
        "4": 1,
    }
    assert [n.attrs["data-act"] for n in slide.find("button", "ex__btn")] == [
        "check",
        "show",
        "reset",
    ]
    for button in slide.find("button", "ex__btn"):
        assert "ui" in button.classes
    assert slide.one(class_="match__key").text() == "Answers: 1-E, 2-A, 3-C, 4-D, 5-B"


def test_gap_fill_chips_and_word_box(deck: Deck) -> None:
    slide = slides_of(deck, "gap_fill")[0]
    gaps = slide.find("button", "gap")
    assert [g.attrs["data-answer"] for g in gaps] == ["takes", "lives", "visit", "are"]
    assert all(g.attrs["data-act"] == "gap" for g in gaps)
    assert [c.attrs["data-word"] for c in slide.find(class_="wordbox__chip")] == [
        "takes",
        "lives",
        "visit",
        "are",
        "goes",
    ]
    assert [g.one(class_="gap__answer").text() for g in gaps] == [
        g.attrs["data-answer"] for g in gaps
    ]
    assert not slides_of(deck, "gap_fill")[1].find(class_="wordbox__chip")


def test_choice_marks_the_right_option(deck: Deck) -> None:
    slide = slides_of(deck, "choice")[0]
    items = slide.find(class_="choice")
    assert [i.attrs["data-answer"] for i in items] == ["1", "0", "1", "0"]
    for item in items:
        options = item.find("button", "choice__option")
        correct = [o for o in options if o.attrs.get("data-correct") == "true"]
        assert len(correct) == 1
        assert correct[0].attrs["data-index"] == item.attrs["data-answer"]
        assert [o.one(class_="choice__letter").text() for o in options][:2] == ["A", "B"]
    assert "choice-list--cols" in slide.one(class_="choice-list").classes


def test_true_false_buttons(deck: Deck) -> None:
    slide = slides_of(deck, "true_false")[0]
    items = slide.find(class_="tf")
    assert [i.attrs["data-answer"] for i in items] == ["true", "false", "true"]
    for item in items:
        buttons = item.find("button", "tf__btn")
        assert [b.attrs["data-value"] for b in buttons] == ["true", "false"]
        assert [b.text() for b in buttons] == ["True", "False"]
        right = [b for b in buttons if b.attrs.get("data-correct")]
        assert [b.attrs["data-value"] for b in right] == [item.attrs["data-answer"]]


def test_rule_box_card_with_speakable_examples(deck: Deck) -> None:
    slide = slides_of(deck, "rule_box")[0]
    assert slide.attrs["data-slide-type"] == "rule"
    box = slide.one(class_="rule-box")
    assert box.one(class_="rule-box__title").text() == "Remember!"
    assert len(box.find(class_="rule-box__point")) == 2
    says = [b.attrs["data-say"] for b in box.find("button", "speak")]
    assert says == ["She takes the bus every day.", "They live near the river."]


def test_reading_highlights_reveal_meanings_and_paragraphs_speak(deck: Deck) -> None:
    slide = slides_of(deck, "reading")[0]
    paragraphs = slide.find("p", "reading__para")
    assert len(paragraphs) == 2
    words = slide.find("button", "reading__word")
    assert [(w.text(), w.attrs["data-meaning"]) for w in words] == [
        ("busy", "full of activity"),
        ("crowded", "too many people"),
        ("quiet", "calm"),
    ]
    assert all(w.attrs["data-act"] == "word" for w in words)
    assert len(slide.find("button", "speak")) == 2
    assert slide.find(class_="reading__tip")
    assert slide.one(class_="ex__instruction").text() == "Life in a big city"
    assert slide.one(class_="ex__sub").text() == "Read the text about city life."


def test_reading_without_a_glossary_marks_words_without_buttons(tmp_path: Path) -> None:
    block = changed(READING, glossary={})
    lesson = make_exercise_lesson()
    lesson.blocks["e_read"] = BlockAdapter.validate_python(block)
    deck = load(export_html(lesson, tmp_path / "d.html", theme="sunny"))
    slide = slides_of(deck, "reading")[0]
    assert not slide.find("button", "reading__word")
    assert [m.text() for m in slide.find("mark", "reading__word")] == ["busy", "crowded", "quiet"]
    assert not slide.find(class_="reading__tip")


def test_questions_are_revealed_one_at_a_time(deck: Deck) -> None:
    slide = slides_of(deck, "question_set")[0]
    questions = slide.find(class_="qset__q")
    assert ["is-revealed" in q.classes for q in questions] == [True, False, False]
    assert [a.text() for a in slide.find(class_="qset__answer")] == [
        "I go to school by bike.",
        "I live in a village.",
    ]
    assert [b.attrs["data-act"] for b in slide.find("button", "ex__btn")] == [
        "next",
        "show",
        "reset",
    ]


def test_answer_keys_go_to_the_teacher_notes(deck: Deck) -> None:
    for kind, needle in (
        ("match", "1-E"),
        ("gap_fill", "Mai takes the bus to school."),
        ("choice", "B. goes"),
        ("true_false", "3. True"),
        ("question_set", "I go to school by bike."),
    ):
        notes = slides_of(deck, kind)[0].one(class_="notes")
        assert "teacher" in notes.classes
        assert "Answer key" in notes.text() and needle in notes.text(), kind
    glossary = slides_of(deck, "reading")[0].one(class_="notes").text()
    assert "busy: full of activity" in glossary
    split = slides_of(deck, "gap_fill")[2].one(class_="notes").text()
    assert "7. Sentence 7: he word7 every day." in split and "Sentence 1" not in split


def test_exercise_buttons_never_advance_the_slide_and_reset_on_entry(deck: Deck) -> None:
    script = deck.root.one("script").text()
    assert "button[data-act]" in script  # one dispatcher for every control
    assert "closest('button, a, .notes, .dots, .ex')" in script  # clicks inside never navigate
    assert "if (entering) resetExercises(shown);" in script
    for action in ("pick", "check", "gap", "option", "tf", "word", "show", "next", "reset"):
        assert f"case '{action}'" in script


def test_print_hides_the_controls_and_shows_the_answers(deck: Deck) -> None:
    css = deck.root.find("style")[0].text()
    printing = css[css.index("@media print") :]
    assert re.search(r"\.ex__bar[^{]*\{\s*display:\s*none\s*!important", printing)
    assert re.search(r"\.gap \.gap__answer\s*\{\s*visibility:\s*visible\s*!important", printing)
    assert re.search(r"\.qset__q\s*\{\s*visibility:\s*visible\s*!important", printing)
    assert re.search(r"\.match__key, \.qset__answer\s*\{\s*display:\s*block", printing)
    assert "[data-correct]" in printing
    assert "pointer-events: none" in printing
    on_screen = css[: css.index("@media print")]
    assert re.search(r"\.match__key, \.qset__answer \{ display: none; \}", on_screen)


def test_exercise_markup_has_no_external_reference(deck: Deck) -> None:
    assert "http://" not in deck.text and "https://" not in deck.text
    assert "url(" not in deck.text.lower().split("<body")[0].split("/* --- theme")[-1]


def test_every_class_of_the_exercise_deck_is_documented(deck: Deck) -> None:
    readme = THEMES_README.read_text(encoding="utf-8")
    emitted: set[str] = set()
    for node in deck.root.walk():
        emitted |= node.classes
    states = {"is-correct", "is-wrong", "is-picked", "is-revealed", "is-done"}
    assert states <= set(re.findall(r"is-[a-z]+", readme))
    undocumented = sorted(c for c in emitted | states if f".{c}" not in readme)
    assert undocumented == []


@pytest.mark.parametrize("theme", ["kids", "storybook", "chalkboard", "candy", "studio", "plain"])
def test_every_theme_renders_the_exercises(theme: str, tmp_path: Path, lesson: Lesson) -> None:
    deck = load(export_html(lesson, tmp_path / f"{theme}.html", theme=theme))
    assert len(deck.root.find("section", "slide", data_slide_type="exercise")) == 8
    assert deck.root.find(class_="rule-box")


def test_user_text_in_exercises_is_escaped(tmp_path: Path) -> None:
    lesson = make_exercise_lesson()
    lesson.blocks["e_tf"] = BlockAdapter.validate_python(
        changed(TRUE_FALSE, items=[{"statement": '<b>bold</b> & "quoted"', "answer": True}])
    )
    text = export_html(lesson, tmp_path / "x.html").read_text(encoding="utf-8")
    assert "<b>bold</b>" not in text
    assert "&lt;b&gt;bold&lt;/b&gt; &amp; &quot;quoted&quot;" in text


# --- PowerPoint --------------------------------------------------------------------------------


def slide_text(slide: Any) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def test_pptx_has_a_readable_slide_per_exercise(tmp_path: Path, lesson: Lesson) -> None:
    path = export_pptx(lesson, tmp_path / "deck.pptx")
    slides = list(Presentation(str(path)).slides)
    texts = [slide_text(slide) for slide in slides]
    notes = [
        slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "" for slide in slides
    ]
    match = next(i for i, t in enumerate(texts) if t.startswith("Match 2"))
    assert "Match the words with their meanings." in texts[match]
    assert "1. busy" in texts[match] and "A. calm, with little noise" in texts[match]
    assert "1-E  busy = full of activity" in notes[match]

    gap = next(i for i, t in enumerate(texts) if "Word box: takes / lives" in t)
    assert "1. Mai ___ the bus to school." in texts[gap]
    assert "1. Mai takes the bus to school." in notes[gap]

    rule = next(i for i, t in enumerate(texts) if t.startswith("Remember!"))
    assert "Use the present simple for habits and facts." in texts[rule]
    assert "e.g. She takes the bus every day." in texts[rule]

    reading = next(i for i, t in enumerate(texts) if t.startswith("Life in a big city"))
    assert "busy city" in texts[reading]
    assert "busy: full of activity" in notes[reading]

    assert any("1. She ___ to school by bus." in t and "B. goes" in t for t in texts)
    assert any("(True / False)" in t for t in texts)
    assert any("3. What do you do in the evening?" in t for t in texts)


def test_pptx_splits_long_exercises(tmp_path: Path, lesson: Lesson) -> None:
    slides = list(Presentation(str(export_pptx(lesson, tmp_path / "d.pptx"))).slides)
    titles = [slide.shapes.title.text_frame.text for slide in slides if slide.shapes.title]
    assert "Fill in the gaps 6 (1/2)" in titles and "Fill in the gaps 6 (2/2)" in titles


# --- Word --------------------------------------------------------------------------------------


def test_docx_puts_the_exercise_with_the_students_and_the_key_with_the_teacher(
    tmp_path: Path, lesson: Lesson
) -> None:
    path = export_docx(lesson, tmp_path / "plan.docx")
    tables = Document(str(path)).tables
    rows = [
        (row.cells[0].text, row.cells[1].text)
        for table in tables
        for row in table.rows
        if "(exercise)" in row.cells[0].text
    ]
    by_title = {
        re.sub(r"\s*AI-added.*$", "", teacher.splitlines()[0]): (teacher, students)
        for teacher, students in rows
    }
    teacher, students = by_title["Match 2 (exercise)"]
    assert "Instructions: Match the words with their meanings." in teacher
    assert "Answer key:" in teacher and "1-E  busy = full of activity" in teacher
    assert "1. busy" in students and "A. calm, with little noise" in students
    assert "Answer key" not in students

    teacher, students = by_title["Fill in the gaps 3 (exercise)"]
    assert "1. Mai takes the bus to school." in teacher
    assert "Word box: takes / lives / visit / are / goes" in students
    assert "Mai ___ the bus" in students

    teacher, students = by_title["Life in a big city (exercise)"]
    assert "busy: full of activity" in teacher and "busy city" in students
    assert "Remember! (exercise)" in by_title


# --- prompt pack -------------------------------------------------------------------------------


def test_prompt_pack_lists_the_exercise_and_keeps_the_key_teacher_only(lesson: Lesson) -> None:
    text = render_lesson_for_prompt(lesson)
    assert "### e_match · match" in text
    assert "- instruction: Match the words with their meanings." in text
    assert "- 1. busy" in text
    assert "- glossary: busy = full of activity" in text
    teacher_only = text[text.index("### e_match") :].split("###")[1]
    assert "Teacher-only:" in teacher_only and "answer key:" in teacher_only
    body = teacher_only.split("Teacher-only:")[0]
    assert "1-E" not in body


# --- the local app -----------------------------------------------------------------------------


def test_the_app_labels_every_exercise_block(lesson: Lesson) -> None:
    from classdeck.app.server import _block_label

    labels = {block_id: _block_label(block) for block_id, block in lesson.blocks.items()}
    assert labels["e_read"] == "Life in a big city"
    assert labels["e_rule"] == "Remember!"
    assert labels["e_match"] == "Match the words with their meanings."
    assert all(labels.values())
