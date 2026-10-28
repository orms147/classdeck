from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as Deck
from pptx.slide import Slide
from pptx.util import Inches, Pt

from classdeck.curriculum import load_curriculum
from classdeck.export import (
    AI_ADDED_BADGE,
    AI_ADDED_NOTE,
    PROVENANCE_LABEL,
    ExportBlocked,
    SourceMapRow,
    build_source_map,
    collect_failures,
    collect_warnings,
    ensure_exportable,
    export_docx,
    export_pptx,
    format_check,
    format_citation,
    origin_text,
)
from classdeck.export.docx import WRAP_UP_FALLBACK
from classdeck.export.pptx import ACCENT
from classdeck.lesson import (
    BlockAdapter,
    CheckResult,
    CheckStatus,
    Citation,
    Lesson,
    Phase,
    Provenance,
    Section,
    SourceDocument,
)
from tests.test_lesson_schema import (
    DOC_ID,
    IPA_HELLO,
    IPA_HI,
    make_lesson,
    make_source_document,
)

CURRICULUM_ROOT = Path(__file__).resolve().parents[1] / "curriculum"
WARNING = CheckResult(name="english.vocab_level", status=CheckStatus.WARN, detail="1 off-list word")
FAILURE = CheckResult(name="english.steps_present", status=CheckStatus.FAIL, detail="no steps")


def _cite(span_id: str) -> dict[str, Any]:
    return {"doc_id": DOC_ID, "page": 6, "span_id": span_id}


def make_export_lesson() -> Lesson:
    """The schema test lesson plus one block of every exported kind and all five phases."""
    lesson = make_lesson()
    extra: list[dict[str, Any]] = [
        {
            "id": "b_04",
            "kind": "activity",
            "content": {
                "type": "chant",
                "title": "Hello chant",
                "skills": ["speaking"],
                "interaction": "whole_class",
                "minutes": 4,
                "instructions": "Chant with the class.",
                "teacher_steps": ["Chant the first line.", "Clap the rhythm."],
                "student_steps": ["Listen.", "Chant and clap."],
                "materials": ["flashcards"],
            },
            "sources": [_cite("p6.s2")],
            "provenance": "extracted",
        },
        {
            "id": "b_05",
            "kind": "dialogue",
            "content": {
                "context": "Ha meets Tien.",
                "lines": [
                    {"speaker": "Ha", "text": "Hello, I'm Ha."},
                    {"speaker": "Tien", "text": "Hi, Ha. I'm Tien."},
                ],
            },
            "sources": [_cite("p6.s3")],
            "provenance": "extracted",
        },
        {
            "id": "b_06",
            "kind": "phonics",
            "content": {
                "sound": "/h/",
                "letters": "h",
                "words": ["hello", "hi"],
                "chant_lines": ["H, h, hello!"],
            },
            "sources": [_cite("p6.s6")],
            "provenance": "extracted",
        },
        {
            "id": "b_09",
            "kind": "image_ref",
            "content": {"caption": "Ha and Tien say hello.", "alt": "Two children waving"},
            "sources": [_cite("p6.s4")],
            "provenance": "extracted",
        },
        {
            "id": "b_10",
            "kind": "activity",
            "content": {
                "type": "listen_point_say",
                "title": "Listen, point and say",
                "skills": ["listening", "speaking"],
                "interaction": "pairs",
                "minutes": 8,
                "teacher_steps": ["Say a word.", "Check the pointing."],
                "student_steps": ["Point to the picture.", "Say the word."],
                "materials": ["a soft ball", "flashcards"],
                "competency_tags": ["digital competence"],
            },
            "provenance": "derived",
            "depends_on": ["b_02"],
        },
        {
            "id": "b_11",
            "kind": "teacher_note",
            "content": {"text": "Praise effort, not only accuracy."},
            "provenance": "generated",
        },
    ]
    for data in extra:
        lesson.blocks[data["id"]] = BlockAdapter.validate_python(data)

    presentation, production, wrap_up = lesson.sections
    presentation.block_ids += ["b_05", "b_09"]
    production.block_ids.append("b_11")
    lesson.sections = [
        Section(
            id="s_warm_up",
            phase=Phase.WARM_UP,
            title="Warm-up: Hello chant",
            minutes=5,
            block_ids=["b_04"],
        ),
        presentation,
        Section(
            id="s_practice",
            phase=Phase.PRACTICE,
            title="Practice: Listen, point and say",
            minutes=12,
            block_ids=["b_06", "b_10"],
        ),
        production,
        wrap_up,
    ]
    return lesson


def write_png(image_root: Path) -> Path:
    span = make_source_document().find_span(6, "p6.s4")
    assert span is not None
    assert span.image_path is not None
    path = image_root / span.image_path
    path.parent.mkdir(parents=True, exist_ok=True)
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 32, 32), False)
    pixmap.clear_with(200)
    pixmap.save(str(path))
    return path


@pytest.fixture
def lesson() -> Lesson:
    return make_export_lesson()


@pytest.fixture
def sources() -> dict[str, SourceDocument]:
    return {DOC_ID: make_source_document()}


@pytest.fixture
def image_root(tmp_path: Path) -> Path:
    root = tmp_path / "workdir"
    write_png(root)
    return root


@pytest.fixture
def deck(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument], image_root: Path
) -> Deck:
    path = export_pptx(lesson, tmp_path / "slides.pptx", sources=sources, image_root=image_root)
    return Presentation(str(path))


# --- slide helpers -----------------------------------------------------------------------------


def shape_texts(slide: Slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def slide_text(slide: Slide) -> str:
    return "\n".join(shape_texts(slide))


def title_of(slide: Slide) -> str:
    return slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""


def notes_of(slide: Slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def slide_titled(deck: Deck, title: str) -> Slide:
    matches = [slide for slide in deck.slides if title_of(slide) == title]
    assert len(matches) == 1, f"{len(matches)} slides titled {title!r}"
    return matches[0]


def shape_named(slide: Slide, name: str) -> Any:
    return next((shape for shape in slide.shapes if shape.name == name), None)


def pictures(slide: Slide) -> list[Any]:
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def table_rows(shape: Any) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in shape.table.rows]


# --- source map --------------------------------------------------------------------------------


def test_format_citation_uses_page_and_paragraph_number() -> None:
    assert format_citation(Citation(doc_id="d", page=6, span_id="p6.s3")) == "p. 6, ¶3"
    assert format_citation(Citation(doc_id="d", page=12, span_id="p12.s10")) == "p. 12, ¶10"


def test_format_citation_falls_back_to_the_span_id() -> None:
    assert format_citation(Citation(doc_id="d", page=6, span_id="img-4")) == "p. 6, img-4"
    assert format_citation(Citation(doc_id="d", page=6, span_id="p6.s3b")) == "p. 6, p6.s3b"


def test_source_map_has_one_row_per_placed_block_in_lesson_order(lesson: Lesson) -> None:
    rows = build_source_map(lesson)
    assert [row.block_id for row in rows] == [block.id for block in lesson.ordered_blocks()]
    assert all(isinstance(row, SourceMapRow) for row in rows)


def test_source_map_rows_carry_provenance_citations_and_quotes(lesson: Lesson) -> None:
    by_id = {row.block_id: row for row in build_source_map(lesson)}
    vocabulary = by_id["b_02"]
    assert vocabulary.kind == "vocabulary"
    assert vocabulary.provenance is Provenance.EXTRACTED
    assert vocabulary.citations == ["p. 6, ¶3"]
    assert vocabulary.quote == "Hello, I'm Ha."
    assert by_id["b_07"].provenance is Provenance.GENERATED
    assert by_id["b_07"].citations == []
    assert by_id["b_07"].quote is None
    assert by_id["b_08"].depends_on == ["b_02", "b_07"]


def test_source_map_quotes_the_span_text_only_when_the_citation_has_none(
    lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    assert {row.block_id: row.quote for row in build_source_map(lesson)}["b_03"] is None
    by_id = {row.block_id: row for row in build_source_map(lesson, sources)}
    assert by_id["b_03"].quote == "Hello, I'm ___. Hi, ___. I'm ___."
    assert by_id["b_02"].quote == "Hello, I'm Ha."


def test_source_map_quote_is_cut_at_eighty_characters(lesson: Lesson) -> None:
    long_text = "abcdefghij" * 20
    document = make_source_document()
    document.pages[0].spans[4].text = long_text
    rows = build_source_map(lesson, {DOC_ID: document})
    quote = {row.block_id: row.quote for row in rows}["b_03"]
    assert quote is not None
    assert len(quote) == 80
    assert long_text.startswith(quote)


def test_source_map_ignores_unresolvable_and_empty_spans(lesson: Lesson) -> None:
    lesson.blocks["b_03"].sources = [Citation(doc_id=DOC_ID, page=6, span_id="p6.s4")]
    rows = build_source_map(lesson, {DOC_ID: make_source_document()})
    assert {row.block_id: row.quote for row in rows}["b_03"] is None
    by_id = {row.block_id: row for row in build_source_map(lesson, {})}
    assert by_id["b_03"].quote is None


def test_origin_text_names_citations_or_dependencies(lesson: Lesson) -> None:
    by_id = {row.block_id: row for row in build_source_map(lesson)}
    assert origin_text(by_id["b_02"]) == "p. 6, ¶3"
    assert origin_text(by_id["b_08"]) == "Derived from b_02, b_07"
    assert origin_text(by_id["b_07"]) == "—"


def test_provenance_labels() -> None:
    assert PROVENANCE_LABEL == {
        Provenance.EXTRACTED: "Source",
        Provenance.DERIVED: "Derived",
        Provenance.GENERATED: "AI-added",
    }
    assert AI_ADDED_BADGE == "AI-added"
    assert AI_ADDED_NOTE == "AI-added — check before teaching"


# --- the export gate ---------------------------------------------------------------------------


def test_collect_checks_scan_blocks_and_the_lesson(lesson: Lesson) -> None:
    lesson.blocks["b_02"].checks.append(WARNING)
    lesson.blocks["b_07"].checks.append(FAILURE)
    lesson.checks.append(WARNING.model_copy(update={"name": "english.timing"}))
    lesson.checks.append(CheckResult(name="grounding.citation_supported", status="stale"))
    assert [(block_id, check.name) for block_id, check in collect_warnings(lesson)] == [
        ("b_02", "english.vocab_level"),
        (None, "english.timing"),
    ]
    assert collect_failures(lesson) == [("b_07", FAILURE)]


def test_ensure_exportable_passes_on_warnings_and_raises_on_failures(lesson: Lesson) -> None:
    lesson.blocks["b_02"].checks.append(WARNING)
    ensure_exportable(lesson)
    lesson.checks.append(FAILURE)
    with pytest.raises(ExportBlocked) as raised:
        ensure_exportable(lesson)
    assert raised.value.failures == [(None, FAILURE)]
    assert isinstance(raised.value, ValueError)
    assert "Lesson — english.steps_present: no steps" in str(raised.value)


def test_format_check() -> None:
    assert format_check("b_02", WARNING) == "b_02 — english.vocab_level: 1 off-list word"
    assert format_check(None, CheckResult(name="x", status="warn")) == "Lesson — x"


@pytest.mark.parametrize("writer", [export_pptx, export_docx])
def test_export_refuses_a_lesson_with_a_failed_check(
    writer: Any, tmp_path: Path, lesson: Lesson
) -> None:
    lesson.blocks["b_07"].checks.append(FAILURE)
    out = tmp_path / "out" / "lesson"
    with pytest.raises(ExportBlocked) as raised:
        writer(lesson, out)
    assert raised.value.failures == [("b_07", FAILURE)]
    assert not out.exists()
    assert not out.parent.exists()


# --- PPTX: structure ---------------------------------------------------------------------------


def test_pptx_is_written_and_reopens_as_16_by_9(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    target = tmp_path / "nested" / "deck" / "slides.pptx"
    path = export_pptx(lesson, target, sources=sources)
    assert path == target
    deck = Presentation(str(path))
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, abs=0.001)
    assert deck.core_properties.title == "Grade 3 · Lesson 1 · Period 1 — Unit 1 · Hello"


def test_slide_count(deck: Deck, lesson: Lesson) -> None:
    # title + dividers + objectives(1) + flashcards(2) + pattern + dialogue + phonics
    # + activities(3) + summary + picture + source map; the teacher note has no slide
    assert len(deck.slides) == 1 + 5 + 11 + 1
    assert len(deck.slides) >= len(lesson.sections) + len(lesson.blocks)


def test_title_slide(deck: Deck, lesson: Lesson) -> None:
    text = slide_text(deck.slides[0])
    assert "Unit 1 · Hello" in text
    assert "Grade 3 · Lesson 1 · Period 1" in text
    assert "ta3.u1" not in text


def test_title_slide_skips_missing_parts(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    lesson.curriculum.lesson_no = None
    lesson.curriculum.period = None
    path = export_pptx(lesson, tmp_path / "slides.pptx", sources=sources)
    texts = shape_texts(Presentation(str(path)).slides[0])
    assert "Grade 3" in texts
    assert not any("ta3.u1" in text for text in texts)


def test_each_section_gets_a_divider_before_its_blocks(deck: Deck, lesson: Lesson) -> None:
    titles = [title_of(slide) for slide in deck.slides]
    assert titles[1] == "Warm-up: Hello chant — 5 min"
    for section in lesson.sections:
        assert f"{section.title} — {section.minutes} min" in titles
    assert titles.index("Production: Name chain — 6 min") < titles.index("Name chain")


def test_slide_text_is_at_least_18_pt(deck: Deck) -> None:
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font.size is not None
                    assert run.font.size >= Pt(18), run.text


# --- PPTX: block layouts -----------------------------------------------------------------------


def test_flashcard_has_word_ipa_meaning_and_a_picture(deck: Deck) -> None:
    slide = slide_titled(deck, "Vocabulary 1/2")
    text = slide_text(slide)
    assert "hello" in text
    assert IPA_HELLO in text
    assert "xin chào" in text
    runs = {
        run.text: run
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    }
    assert runs["hello"].font.size >= Pt(40)
    assert runs[IPA_HELLO].font.size < runs["hello"].font.size
    assert runs["xin chào"].font.size < runs[IPA_HELLO].font.size
    assert len(pictures(slide)) == 1
    picture = pictures(slide)[0]
    assert picture.left > deck.slide_width / 2
    assert picture._element.nvPicPr.cNvPr.get("descr") == "hello"
    assert "Example: Hello, I'm Ha." in notes_of(slide)


def test_flashcard_without_a_picture_citation_has_no_picture(deck: Deck) -> None:
    slide = slide_titled(deck, "Vocabulary 2/2")
    assert "\n".join(["hi", IPA_HI, "chào"]) in shape_texts(slide)
    assert pictures(slide) == []


def test_flashcard_picture_needs_sources_and_a_file_under_image_root(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument], image_root: Path
) -> None:
    def vocabulary_pictures(**kwargs: Any) -> int:
        path = export_pptx(lesson, tmp_path / "x.pptx", **kwargs)
        return len(pictures(slide_titled(Presentation(str(path)), "Vocabulary 1/2")))

    assert vocabulary_pictures(sources=sources, image_root=image_root) == 1
    assert vocabulary_pictures(sources=sources) == 0
    assert vocabulary_pictures(image_root=image_root) == 0
    assert vocabulary_pictures(sources=sources, image_root=tmp_path / "empty") == 0


def test_image_path_cannot_escape_image_root(
    tmp_path: Path, lesson: Lesson, image_root: Path
) -> None:
    outside = tmp_path / "outside.png"
    outside.write_bytes(write_png(tmp_path / "scratch").read_bytes())
    document = make_source_document()
    document.pages[0].spans[3].image_path = "../outside.png"
    path = export_pptx(
        lesson, tmp_path / "x.pptx", sources={DOC_ID: document}, image_root=image_root
    )
    assert all(pictures(slide) == [] for slide in Presentation(str(path)).slides)


def test_image_ref_slide_shows_the_cited_picture_and_caption(deck: Deck) -> None:
    slide = slide_titled(deck, "Picture")
    assert len(pictures(slide)) == 1
    assert "Ha and Tien say hello." in slide_text(slide)
    assert "Two children waving" in notes_of(slide)


def test_image_ref_slide_without_the_file_shows_the_alt_text(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    path = export_pptx(lesson, tmp_path / "x.pptx", sources=sources)
    slide = slide_titled(Presentation(str(path)), "Picture")
    assert pictures(slide) == []
    assert "Picture not available: Two children waving" in slide_text(slide)


def test_pattern_slide_highlights_the_slot(deck: Deck) -> None:
    slide = slide_titled(deck, "Sentence pattern")
    paragraphs = [
        paragraph
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]
    pattern = next(p for p in paragraphs if p.text == "Hello, I'm ___.")
    slot = next(run for run in pattern.runs if run.text == "___")
    assert slot.font.bold is True
    assert slot.font.color.rgb == ACCENT
    assert all(not run.font.bold for run in pattern.runs if run.text != "___")
    text = slide_text(slide)
    assert "Introduce yourself" in text
    assert "Hello, I'm Ha." in text


def test_dialogue_slide_has_speaker_labels_and_context(deck: Deck) -> None:
    slide = slide_titled(deck, "Dialogue")
    text = slide_text(slide)
    assert "Ha: Hello, I'm Ha." in text
    assert "Tien: Hi, Ha. I'm Tien." in text
    assert "Ha meets Tien." in text


def test_long_dialogues_continue_on_another_slide(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    dialogue = lesson.blocks["b_05"]
    dialogue.content.lines *= 4  # type: ignore[union-attr]
    path = export_pptx(lesson, tmp_path / "x.pptx", sources=sources)
    deck = Presentation(str(path))
    assert "Dialogue (1/2)" in [title_of(slide) for slide in deck.slides]
    assert "Dialogue (2/2)" in [title_of(slide) for slide in deck.slides]


def test_phonics_slide(deck: Deck) -> None:
    slide = slide_titled(deck, "Phonics")
    text = slide_text(slide)
    assert "/h/" in text
    assert "hello  ·  hi" in text
    assert "H, h, hello!" in text


def test_activity_slide_shows_pupil_steps_and_keeps_teacher_steps_in_the_notes(deck: Deck) -> None:
    slide = slide_titled(deck, "Name chain")
    text = slide_text(slide)
    assert "Pass the ball." in text
    assert "Pass it on." in text
    assert "Model the chain with two pupils." not in text
    notes = notes_of(slide)
    assert "1. Model the chain with two pupils." in notes
    assert "2. Start the ball." in notes
    assert "Materials: a soft ball" in notes
    assert "whole class" in notes
    assert "5 min" in notes
    assert "If the class is large, run two chains in parallel." in notes


def test_summary_and_objective_slides_list_bullets(deck: Deck) -> None:
    assert "hello, hi" in slide_text(slide_titled(deck, "Summary"))
    assert "Pupils can greet each other and say their names." in slide_text(
        slide_titled(deck, "Objectives")
    )


def test_teacher_notes_have_no_slide_and_join_the_previous_slides_notes(deck: Deck) -> None:
    assert all("Praise effort" not in slide_text(slide) for slide in deck.slides)
    notes = notes_of(slide_titled(deck, "Name chain"))
    assert "Teacher note (AI-added): Praise effort, not only accuracy." in notes


# --- PPTX: labelling ---------------------------------------------------------------------------


def test_extracted_slides_cite_their_source_in_a_bottom_left_footer(deck: Deck) -> None:
    for title in ("Vocabulary 1/2", "Vocabulary 2/2"):
        footer = shape_named(slide_titled(deck, title), "Footer")
        assert footer.text_frame.text == "b_02 · Source: p. 6, ¶3"
        assert footer.left < Inches(1)
        assert footer.top > deck.slide_height / 2
    footer = shape_named(slide_titled(deck, "Sentence pattern"), "Footer")
    assert footer.text_frame.text == "b_03 · Source: p. 6, ¶6"


def test_derived_slides_name_what_they_derive_from(deck: Deck) -> None:
    summary = shape_named(slide_titled(deck, "Summary"), "Footer")
    assert summary.text_frame.text == "b_08 · Derived from b_02, b_07"
    objectives = shape_named(slide_titled(deck, "Objectives"), "Footer")
    assert objectives.text_frame.text == "b_01 · Derived from ta3.u1.o1"
    practice = shape_named(slide_titled(deck, "Listen, point and say"), "Footer")
    assert practice.text_frame.text == "b_10 · Derived from b_02"


def test_generated_slides_have_a_badge_and_a_speaker_note(deck: Deck) -> None:
    slide = slide_titled(deck, "Name chain")
    badge = shape_named(slide, "AI-added badge")
    assert badge.text_frame.text == AI_ADDED_BADGE
    assert badge.left > deck.slide_width / 2
    assert badge.top < Inches(1)
    assert AI_ADDED_NOTE in notes_of(slide)
    assert shape_named(slide, "Footer").text_frame.text == "b_07"


def test_grounded_slides_have_no_badge_and_no_ai_note(deck: Deck) -> None:
    for title in ("Vocabulary 1/2", "Sentence pattern", "Dialogue", "Summary", "Hello chant"):
        slide = slide_titled(deck, title)
        assert shape_named(slide, "AI-added badge") is None
        assert AI_ADDED_NOTE not in notes_of(slide)


def test_every_flashcard_of_a_generated_block_is_badged(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    lesson.blocks["b_02"].provenance = Provenance.GENERATED
    path = export_pptx(lesson, tmp_path / "x.pptx", sources=sources)
    deck = Presentation(str(path))
    for title in ("Vocabulary 1/2", "Vocabulary 2/2"):
        slide = slide_titled(deck, title)
        assert shape_named(slide, "AI-added badge") is not None
        assert AI_ADDED_NOTE in notes_of(slide)


# --- PPTX: notes and source map ----------------------------------------------------------------


def test_source_map_slide_lists_every_block(deck: Deck, lesson: Lesson) -> None:
    slide = slide_titled(deck, "Source map")
    assert title_of(deck.slides[-1]) == "Source map"
    table = shape_named(slide, "Source map table")
    rows = table_rows(table)
    assert rows[0] == ["Block", "Kind", "Provenance", "Citation", "Quote"]
    assert len(rows) == len(lesson.ordered_blocks()) + 1
    by_block = {row[0]: row for row in rows[1:]}
    assert by_block["b_02"] == [
        "b_02",
        "vocabulary",
        "Source",
        "p. 6, ¶3",
        "Hello, I'm Ha.",
    ]
    assert by_block["b_07"][2:] == ["AI-added", "—", ""]
    assert by_block["b_08"][2:4] == ["Derived", "Derived from b_02, b_07"]


def test_source_map_quotes_come_from_the_sources_when_given(deck: Deck) -> None:
    rows = table_rows(shape_named(slide_titled(deck, "Source map"), "Source map table"))
    assert {row[0]: row[4] for row in rows}["b_03"] == "Hello, I'm ___. Hi, ___. I'm ___."


def test_source_map_is_split_at_twelve_rows_per_slide(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    for number in range(3):
        block_id = f"b_n{number}"
        lesson.blocks[block_id] = BlockAdapter.validate_python(
            {"id": block_id, "kind": "teacher_note", "content": {"text": "Smile."}}
        )
        lesson.sections[-1].block_ids.append(block_id)
    assert len(lesson.ordered_blocks()) == 14
    path = export_pptx(lesson, tmp_path / "x.pptx", sources=sources)
    deck = Presentation(str(path))
    first = slide_titled(deck, "Source map (1/2)")
    second = slide_titled(deck, "Source map (2/2)")
    assert len(table_rows(shape_named(first, "Source map table"))) == 13
    assert len(table_rows(shape_named(second, "Source map table"))) == 3


def test_notes_slide_appears_only_when_there_are_warnings(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument], deck: Deck
) -> None:
    assert "Notes" not in [title_of(slide) for slide in deck.slides]

    lesson.blocks["b_02"].checks.append(WARNING)
    lesson.checks.append(WARNING.model_copy(update={"name": "english.timing", "detail": "38 min"}))
    path = export_pptx(lesson, tmp_path / "warn.pptx", sources=sources)
    warned = Presentation(str(path))
    titles = [title_of(slide) for slide in warned.slides]
    assert titles.count("Notes") == 1
    assert titles.index("Notes") < titles.index("Source map")
    text = slide_text(slide_titled(warned, "Notes"))
    assert "b_02 — english.vocab_level: 1 off-list word" in text
    assert "Lesson — english.timing: 38 min" in text


# --- DOCX helpers ------------------------------------------------------------------------------


@pytest.fixture
def plan(tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]) -> Any:
    return Document(str(export_docx(lesson, tmp_path / "plan.docx", sources=sources)))


def headings(document: Any) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name.startswith(("Heading", "Title"))]


def paragraphs_under(document: Any, heading: str) -> list[Paragraph]:
    found: list[Paragraph] = []
    inside = False
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading 1"):
            inside = paragraph.text == heading
        elif inside:
            found.append(paragraph)
    return found


def procedure_tables(document: Any) -> list[Any]:
    return [
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["Teacher's activities", "Students' activities"]
    ]


def row_with(document: Any, needle: str, *, column: int = 0) -> Any:
    """The first procedure row whose cell in ``column`` contains ``needle``."""
    for table in procedure_tables(document):
        for row in table.rows[1:]:
            if needle in row.cells[column].text:
                return row
    raise AssertionError(f"no procedure row contains {needle!r}")


def iter_cells(table: Any) -> list[Any]:
    return [cell for row in table.rows for cell in row.cells]


def all_runs(document: Any) -> list[Any]:
    """Every run in body paragraphs and table cells, in document order, merged cells once."""
    runs = {run._r: run for paragraph in document.paragraphs for run in paragraph.runs}
    for table in document.tables:
        for cell in iter_cells(table):
            runs.update({run._r: run for p in cell.paragraphs for run in p.runs})
    return list(runs.values())


# --- DOCX --------------------------------------------------------------------------------------


def test_docx_is_written_and_reopens(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    target = tmp_path / "nested" / "plan.docx"
    path = export_docx(lesson, target, sources=sources)
    assert path == target
    document = Document(str(path))
    assert document.core_properties.title == "Grade 3 · Lesson 1 · Period 1 — Unit 1 · Hello"
    assert "Unit 1 · Hello" in headings(document)


def test_docx_uses_arial_12_so_vietnamese_diacritics_render(plan: Any) -> None:
    normal = plan.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(12)
    assert plan.styles["Heading 1"].font.name == "Arial"


def test_docx_has_the_four_parts_in_order(plan: Any) -> None:
    parts = [text for text in headings(plan) if re.match(r"(I|II|III|IV)\. ", text)]
    assert parts == [
        "I. Objectives",
        "II. Teaching aids",
        "III. Procedure",
        "IV. Wrap-up / Homework",
    ]
    assert headings(plan)[-1] == "Appendix — Source map"
    assert "Notes" not in headings(plan)


def test_docx_header_block(plan: Any) -> None:
    cells = [cell.text for row in plan.tables[0].rows for cell in row.cells]
    assert cells == [
        "Unit: 1 · Hello",
        "Grade: 3",
        "Lesson: 1",
        "Period: 1",
        "Date: ______",
        "Class: ______",
    ]


def test_docx_objectives_list_the_objective_blocks_with_outcome_ids(plan: Any) -> None:
    texts = [p.text for p in paragraphs_under(plan, "I. Objectives")]
    assert texts == ["Pupils can greet each other and say their names. (ta3.u1.o1)"]


def test_docx_expands_curriculum_outcomes_when_a_curriculum_is_given(
    tmp_path: Path, lesson: Lesson
) -> None:
    curriculum = load_curriculum(CURRICULUM_ROOT, "gdpt-2018", 3)
    path = export_docx(lesson, tmp_path / "plan.docx", curriculum=curriculum)
    document = Document(str(path))
    texts = [p.text for p in paragraphs_under(document, "I. Objectives")]
    assert "Curriculum outcomes:" in texts
    assert "ta3.u1.o1 — Greet and respond to greetings using Hello / Hi." in texts
    assert document.tables[0].cell(0, 0).text == "Unit: 1 · Hello"

    plain = Document(str(export_docx(lesson, tmp_path / "plain.docx")))
    assert "Curriculum outcomes:" not in [p.text for p in plain.paragraphs]


def test_docx_unknown_outcome_ids_are_reported_not_fatal(tmp_path: Path, lesson: Lesson) -> None:
    lesson.curriculum.outcome_ids = ["ta3.u9.o9"]
    curriculum = load_curriculum(CURRICULUM_ROOT, "gdpt-2018", 3)
    document = Document(str(export_docx(lesson, tmp_path / "plan.docx", curriculum=curriculum)))
    texts = [p.text for p in paragraphs_under(document, "I. Objectives")]
    assert "ta3.u9.o9 — (not found in the curriculum)" in texts


def test_docx_teaching_aids_are_the_sorted_materials_plus_slides(plan: Any) -> None:
    texts = [p.text for p in paragraphs_under(plan, "II. Teaching aids")]
    assert texts == ["a soft ball", "flashcards", "slides"]


def test_docx_procedure_has_one_two_column_table_per_phase(plan: Any, lesson: Lesson) -> None:
    tables = procedure_tables(plan)
    assert len(tables) == 5
    for table in tables:
        header = table.rows[0].cells
        assert [cell.text for cell in header] == ["Teacher's activities", "Students' activities"]
        assert all(run.bold for cell in header for run in cell.paragraphs[0].runs)
    sections = headings(plan)
    for section in lesson.sections:
        assert f"{section.title} ({section.minutes} min)" in sections


def test_docx_presentation_table_has_a_row_per_teaching_block(plan: Any) -> None:
    table = next(
        t for t in procedure_tables(plan) if "Present vocabulary" in t.rows[1].cells[0].text
    )
    assert len(table.rows) == 1 + 3  # vocabulary, pattern, dialogue; objective and picture skipped
    vocabulary, pattern, dialogue = (row.cells for row in table.rows[1:])
    assert f"hello {IPA_HELLO} — xin chào" in vocabulary[0].text
    assert "hi /haɪ/ — chào" in vocabulary[0].text  # noqa: RUF001
    assert vocabulary[1].text == "Listen, repeat, point"
    assert "Sentence pattern: Hello, I'm ___." in pattern[0].text
    assert "Example: Hello, I'm Ha." in pattern[0].text
    assert pattern[1].text == "Repeat and practise the pattern"
    assert "Ha: Hello, I'm Ha." in dialogue[0].text
    assert dialogue[1].text == "Listen and repeat"


def test_docx_activity_row_numbers_teacher_and_student_steps(plan: Any) -> None:
    row = row_with(plan, "Name chain")
    teacher, students = row.cells
    assert "1. Model the chain with two pupils." in teacher.text
    assert "2. Start the ball." in teacher.text
    assert "Name chain (5 min, whole class)" in teacher.text
    assert students.text.splitlines() == [
        "1. Catch the ball, say 'Hello, I'm (name).'",
        "2. Pass it on.",
    ]


def test_docx_phonics_and_summary_rows(plan: Any) -> None:
    phonics = row_with(plan, "Sound: /h/")
    assert "Words: hello, hi" in phonics.cells[0].text
    assert "H, h, hello!" in phonics.cells[0].text
    assert phonics.cells[1].text == "Listen and repeat the sound"
    summary = row_with(plan, "Recall with the teacher", column=1)
    first, second = summary.cells[0].text.splitlines()
    assert first == "hello, hi"
    assert re.fullmatch(r"Hello, I'm ___\.\[\d+\]", second)


def test_docx_competences_line_follows_the_activity_and_is_highlighted(plan: Any) -> None:
    table = next(t for t in procedure_tables(plan) if "Name chain" in t.rows[1].cells[0].text)
    assert len(table.rows) == 3
    line = table.rows[2].cells[0]
    assert line.text == "Competences: communication, cooperation"
    run = line.paragraphs[0].runs[0]
    assert run.bold is True
    assert run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    highlighted = [
        r.text for r in all_runs(plan) if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert highlighted == [
        "Competences: digital competence",
        "Competences: communication, cooperation",
    ]


def test_docx_rows_end_with_a_superscript_marker_into_the_source_map(
    plan: Any, lesson: Lesson
) -> None:
    numbers = {row.block_id: n for n, row in enumerate(build_source_map(lesson), start=1)}
    rows = [
        row
        for table in procedure_tables(plan)
        for row in table.rows[1:]
        if not row.cells[0].text.startswith("Competences:")
    ]
    assert len(rows) == 8
    for row in rows:
        last = row.cells[0].paragraphs[-1].runs[-1]
        assert re.fullmatch(r"\[\d+\]", last.text)
        assert last.font.superscript is True
    assert row_with(plan, "Present vocabulary").cells[0].paragraphs[-1].runs[-1].text == (
        f"[{numbers['b_02']}]"
    )
    assert row_with(plan, "Name chain").cells[0].paragraphs[-1].runs[-1].text == (
        f"[{numbers['b_07']}]"
    )


def test_docx_marks_only_generated_rows_with_the_ai_added_badge(plan: Any) -> None:
    generated = row_with(plan, "Name chain").cells[0].paragraphs[-1].runs
    badge = next(run for run in generated if run.text == AI_ADDED_BADGE)
    assert badge.bold is True
    assert badge.font.highlight_color is not None
    assert generated[-1].text.startswith("[")
    grounded = row_with(plan, "Present vocabulary").cells[0].paragraphs[-1].runs
    assert all(run.text != AI_ADDED_BADGE for run in grounded)


def test_docx_teacher_notes_are_italic_paragraphs_under_the_table(
    plan: Any, lesson: Lesson
) -> None:
    numbers = {row.block_id: n for n, row in enumerate(build_source_map(lesson), start=1)}
    notes = {
        p.runs[0].text: p
        for p in plan.paragraphs
        if p.runs and p.runs[0].italic and not p.style.name.startswith("Heading")
    }
    block_note = notes["Praise effort, not only accuracy."]
    assert [run.text for run in block_note.runs] == [
        "Praise effort, not only accuracy.",
        " ",
        AI_ADDED_BADGE,
        f"[{numbers['b_11']}]",
    ]
    field_note = notes["If the class is large, run two chains in parallel."]
    assert field_note.runs[-1].text == f"[{numbers['b_07']}]"
    assert all(
        "Praise effort" not in cell.text for table in plan.tables for cell in iter_cells(table)
    )


def test_docx_wrap_up_shows_the_wrap_up_blocks(plan: Any) -> None:
    texts = [p.text for p in paragraphs_under(plan, "IV. Wrap-up / Homework")]
    assert texts == ["Wrap-up (2 min)"]
    last = procedure_tables(plan)[-1]
    assert last.rows[1].cells[1].text == "Recall with the teacher"


def test_docx_wrap_up_falls_back_to_a_homework_sentence(tmp_path: Path, lesson: Lesson) -> None:
    lesson.sections = [s for s in lesson.sections if s.phase is not Phase.WRAP_UP]
    document = Document(str(export_docx(lesson, tmp_path / "plan.docx")))
    texts = [p.text for p in paragraphs_under(document, "IV. Wrap-up / Homework")]
    assert texts == [WRAP_UP_FALLBACK]
    assert WRAP_UP_FALLBACK == "Review the new words and the pattern at home."
    assert all(
        "Recall with the teacher" not in cell.text
        for table in document.tables
        for cell in iter_cells(table)
    )


def test_docx_notes_section_lists_warnings_only_when_there_are_some(
    tmp_path: Path, lesson: Lesson, sources: dict[str, SourceDocument]
) -> None:
    lesson.blocks["b_02"].checks.append(WARNING)
    lesson.checks.append(WARNING.model_copy(update={"name": "english.timing", "detail": "38 min"}))
    document = Document(str(export_docx(lesson, tmp_path / "plan.docx", sources=sources)))
    assert headings(document).index("Notes") < headings(document).index("Appendix — Source map")
    assert [p.text for p in paragraphs_under(document, "Notes")] == [
        "b_02 — english.vocab_level: 1 off-list word",
        "Lesson — english.timing: 38 min",
    ]


def test_docx_appendix_is_the_source_map(plan: Any, lesson: Lesson) -> None:
    table = plan.tables[-1]
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows[0] == ["n", "Block", "Kind", "Provenance", "Citation", "Quote"]
    assert len(rows) == len(lesson.ordered_blocks()) + 1
    assert [row[0] for row in rows[1:]] == [str(n) for n in range(1, len(rows))]
    assert [row[1] for row in rows[1:]] == [block.id for block in lesson.ordered_blocks()]
    by_block = {row[1]: row for row in rows[1:]}
    assert by_block["b_02"][2:] == ["vocabulary", "Source", "p. 6, ¶3", "Hello, I'm Ha."]
    assert by_block["b_03"][5] == "Hello, I'm ___. Hi, ___. I'm ___."
    assert by_block["b_07"][3:5] == ["AI-added", "—"]
    assert by_block["b_08"][3:5] == ["Derived", "Derived from b_02, b_07"]


def test_docx_appendix_without_sources_keeps_only_given_quotes(
    tmp_path: Path, lesson: Lesson
) -> None:
    document = Document(str(export_docx(lesson, tmp_path / "plan.docx")))
    rows = [[cell.text for cell in row.cells] for row in document.tables[-1].rows[1:]]
    by_block = {row[1]: row for row in rows}
    assert by_block["b_02"][5] == "Hello, I'm Ha."
    assert by_block["b_03"][5] == ""


def test_docx_marker_numbers_match_the_appendix_rows(plan: Any) -> None:
    appendix = {row.cells[1].text: row.cells[0].text for row in plan.tables[-1].rows[1:]}
    marker = row_with(plan, "Sentence pattern").cells[0].paragraphs[-1].runs[-1].text
    assert marker == f"[{appendix['b_03']}]"
