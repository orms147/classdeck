"""The teacher's plan as reference: the lesson's section, cleaned and capped."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from classdeck.planning.reference import REFERENCE_CHARS, clean_lines, reference_for_lesson


def build_plan(path: Path) -> Path:
    """A made-up two-lesson plan in the layout of a Vietnamese unit plan: one header table per
    period, objectives paragraphs, then a procedure table with merged cells."""
    document = docx.Document()
    for number, name, topic in (
        (1, "Getting started", "a town park"),
        (2, "A closer look 1", "jobs"),
    ):
        header = document.add_table(rows=1, cols=2)
        header.cell(0, 0).text = "Date of planning: 1 / 9 / 2023\nDate of teaching: ....\nWEEK: 1"
        header.cell(
            0, 1
        ).text = f"Period {number}: UNIT 1\nLesson {number}: {name.upper()}\n{topic}"
        document.add_paragraph("I. OBJECTIVES: pupils will talk about the topic;")
        document.add_paragraph("* New words: park; bench; playground;")
        document.add_paragraph(
            f"Digital Competence (1.1): pupils search the web for pictures of {topic}."
        )
        document.add_paragraph("Digital evidence: a screenshot.")
        document.add_paragraph("Objectives for students with disabilities")
        document.add_paragraph("By the end of the lesson, students can:")
        document.add_paragraph("- copy two words")
        document.add_paragraph("II. TEACHING AIDS:")
        document.add_paragraph("- Hoclieu.vn")
        document.add_paragraph("- Projector")
        document.add_paragraph("III. PROCEDURE:")
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = f"1. WARM UP (5 min) of lesson {number}"
        table.cell(1, 0).text = (
            f"- Ask pupils: Where do you play? (lesson {number})\n"
            "- Use Google Images to show a park.\n"
            "For students with disabilities\n"
            "- Student points at one picture."
        )
        table.cell(1, 1).text = "- T_Ss\n* Key: 1. park"
        table.cell(2, 0).text = "Wrap-up: recap the words."
        table.cell(2, 1).text = "Homework: learn the words."
        document.add_paragraph("========================================")
    document.save(str(path))
    return path


@pytest.fixture
def plan(tmp_path: Path) -> Path:
    return build_plan(tmp_path / "plan.docx")


def test_the_lessons_own_section_is_returned(plan: Path) -> None:
    first = reference_for_lesson(plan, 1)
    assert "Lesson 1: GETTING STARTED" in first
    assert "Ask pupils: Where do you play? (lesson 1)" in first
    assert "lesson 2" not in first and "A CLOSER LOOK" not in first
    second = reference_for_lesson(plan, 2)
    assert "Lesson 2: A CLOSER LOOK 1" in second and "(lesson 2)" in second
    assert "(lesson 1)" not in second


def test_digital_special_needs_and_date_lines_are_dropped(plan: Path) -> None:
    text = reference_for_lesson(plan, 1)
    lowered = text.lower()
    for word in ("digital", "google", "hoclieu", "disabilit", "date of", "week:", "copy two words"):
        assert word not in lowered, word
    assert "student points" not in lowered  # the bullet under a special-needs line
    assert "* Key: 1. park" in text  # content cells of the procedure table stay
    assert "New words: park; bench; playground;" in text  # the objectives block keeps the words
    assert "pupils will talk about the topic" not in text  # and drops the rest of it
    assert "III. PROCEDURE:" in text and "1. WARM UP (5 min) of lesson 1" in text


def test_a_lesson_the_plan_does_not_have_gives_nothing(plan: Path) -> None:
    assert reference_for_lesson(plan, 5) == ""


def test_the_reference_is_capped_at_a_line_boundary(tmp_path: Path) -> None:
    document = docx.Document()
    header = document.add_table(rows=1, cols=2)
    header.cell(0, 0).text = "Date of planning: x"
    header.cell(0, 1).text = "Lesson 1: Getting started"
    for number in range(400):
        document.add_paragraph(f"Step {number}: do something with the class.")
    path = tmp_path / "long.docx"
    document.save(str(path))
    text = reference_for_lesson(path, 1)
    assert len(text) <= REFERENCE_CHARS and text.splitlines()[-1].startswith("Step ")
    assert len(reference_for_lesson(path, 1, limit=500)) <= 500


def test_clean_lines_resets_special_needs_skipping_after_a_normal_line() -> None:
    lines = clean_lines(
        "For students with disabilities\n- copy a word\nBack to the plan.\n- a normal bullet"
    )
    assert lines == ["Back to the plan.", "- a normal bullet"]
