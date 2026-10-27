"""pagekit: render a unit *spec* into textbook-style PDF, DOCX and PPTX fixtures.

Content writers describe a unit as data; this module lays it out the way a primary English
course book does, so ``classdeck.ingest`` and ``classdeck.planning`` see realistic pages (a
heading, numbered activity lines, pictures with a bounding box, dialogue and pattern lines)
without any copyrighted material. Pictures are drawn on the fly with PyMuPDF primitives and the
output is deterministic: one spec always renders to byte-identical files.

Spec format (JSON, one file per unit in ``units/<id>.json``; the id is the file stem)
-------------------------------------------------------------------------------------

    {
      "id": "gs3-unit1", "grade": 3, "unit": 1, "title": "Hello",
      "series": "classdeck-sample",            optional; goes into the file metadata
      "layout": { "lesson_line": true },       optional; false drops the "Lesson L" line
      "lessons": [                             one to three, numbered 1.. in order
        { "no": 1,
          "headings": ["Unit 1  Hello", "..."],  optional; the two page headings verbatim
          "activities": [                      exactly six, numbered 1..6 in order
            { "no": 1, "instruction": "Look, listen and repeat.", "type": "look_listen_repeat",
              "picture": { "scene": "two children greeting at the school gate", "style": "scene" },
              "lines": ["Ha: Hello, I'm Ha.", "Tien: Hi, Ha. I'm Tien."] },
            { "no": 2, "instruction": "Listen, point and say.", "type": "listen_point_say",
              "pattern": "Hello, I'm ___.",
              "items": [ { "word": "Ha", "picture": { "scene": "girl waving", "style": "card" } },
                         { "word": "Tien" } ],
              "lines": ["Hello, I'm ___.", "Hi, ___. I'm ___."] },
            { "no": 3, "instruction": "Let's talk.", "type": "lets_talk",
              "lines": ["Hello. / Hi.", "I'm (name)."] },
            { "no": 4, "instruction": "Listen and tick.", "type": "listen_and_do",
              "options": ["a. Hello, I'm Lan.", "b. Hi, I'm Duc."] },
            { "no": 5, "instruction": "Look, complete and read.", "type": "look_complete_read",
              "lines": ["Hello, I'm _____.", "Hi, _____. I'm _____."] },
            { "no": 6, "instruction": "Let's sing.", "type": "song",
              "lines": ["Hello, hello! I'm Ha.", "Hello, hello! I'm Tien."] }
          ] },
        { "no": 3, "activities": [
            { "no": 1, "instruction": "Listen and repeat.", "type": "phonics",
              "sound": "/h/", "letters": "h", "words": ["hello", "hi"] },
            { "no": 4, "instruction": "Read and match.", "type": "read_and_match",
              "pairs": [["Hello, I'm Ha.", "Hi, Ha. I'm Tien."], ["Goodbye, Tien.", "Bye, Ha."]] },
            "... the other four ..."
        ] }
      ]
    }

Activity types: ``look_listen_repeat, listen_point_say, lets_talk, listen_and_do,
look_complete_read, song, game, phonics, chant, read_and_match, lets_write, project``.
Every content field is optional and any type may carry any of them. What is present renders
in a fixed order: instruction, pattern, picture, items (a row of picture cards with the words
under them), letters and sound, words, options, pairs, lines. An item without a picture gets a
card drawn from its word. ``picture.shapes`` (``face, sun, house, ball, tree``) picks the
drawing; otherwise it is inferred from the scene text. Text is Latin-1 (the built-in PDF fonts
cover nothing else) and no content line may look like an instruction (``^\\d{1,2}[.)]\\s``).

Layout: two pages per lesson, activities 1-3 on the first and 4-6 on the second. Every page
carries the heading ``Unit N  Title`` (20 pt bold); the first page of a lesson adds a
``Lesson L`` line (16 pt bold, a heading to the PDF and DOCX readers, a paragraph in PPTX).
Instructions render as ``N. Instruction.`` (12 pt bold) and are what ``segment_document``
keys on. Short options and words share one line separated by three spaces (ingest collapses
them to one space); longer ones wrap to one per line. Pairs render as two columns, the right
one shifted by one row. A page that does not fit raises ``ValueError``.

API: ``load_specs``, ``validate_spec``, ``plan_pages``, ``text_sequence``, ``render_unit``.
"""

from __future__ import annotations

import io
import json
import math
import re
import zipfile
import zlib
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Annotated, Literal, get_args

import pymupdf
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.util import Inches as PptxInches
from pptx.util import Pt
from pydantic import BaseModel, ConfigDict, Field, ValidationInfo, field_validator, model_validator

FORMATS: tuple[str, ...] = ("pdf", "docx", "pptx")
AUTHOR = "ClassDeck public fixtures"
FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)
PDF_DATE = "D:20260101000000Z"
ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)

INSTRUCTION_RE = re.compile(r"^\s*\d{1,2}[.)]\s")
"""What the pipeline reads as an activity instruction; no content line may match it."""

ActivityType = Literal[
    "look_listen_repeat",
    "listen_point_say",
    "lets_talk",
    "listen_and_do",
    "look_complete_read",
    "song",
    "game",
    "phonics",
    "chant",
    "read_and_match",
    "lets_write",
    "project",
]
ACTIVITY_TYPES: frozenset[str] = frozenset(get_args(ActivityType))
ShapeName = Literal["face", "sun", "house", "ball", "tree"]
SHAPES: tuple[str, ...] = get_args(ShapeName)

GAP = "   "
"""Separates texts that share one line; ingest collapses it to a single space."""

# --- page geometry (PDF points; the other formats follow the same proportions) ----------------

PAGE_WIDTH, PAGE_HEIGHT = 595.0, 842.0
MARGIN = 60.0
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN
BOTTOM_LIMIT = PAGE_HEIGHT - 42.0
HEADING_BASELINE = 70.0
HEADING_STEP = 34.0  # heading to lesson line; wide enough that MuPDF keeps them as two blocks
LESSON_STEP = 24.0
HEADING_SIZE, LESSON_SIZE, BODY_SIZE, CAPTION_SIZE, CARD_WORD_SIZE = 20.0, 16.0, 12.0, 7.0, 9.0
LINE_STEP = 22.0  # body baselines; below 20 pt MuPDF merges neighbouring lines into one block
INSTRUCTION_GAP = 6.0
PICTURE_RISE = 9.0  # a picture's top sits this far above the baseline it takes the place of
PICTURE_DROP = 20.0  # from a picture's bottom to the next baseline
CARD_DROP = 16.0
SCENE_SIZE = (240.0, 120.0)
CARD_SIZE = (72.0, 72.0)
CARD_GAP = 16.0
CARDS_PER_ROW = 5
PAIR_COLUMN_X = 250.0  # where the right column of a "Read and match" row starts
PICTURE_DPI = 144
BODY_FONT, BOLD_FONT = "helv", "hebo"

SLIDE_WIDTH, SLIDE_HEIGHT, SLIDE_MARGIN, SLIDE_LINE = 8.27, 11.69, 0.7, 0.32  # inches, A4

# --- colours (RGB in 0..1) --------------------------------------------------------------------

INK = (0.20, 0.20, 0.25)
INK_SOFT = (0.45, 0.45, 0.50)
WHITE = (1.0, 1.0, 1.0)
SKIN = (1.0, 0.87, 0.72)
SUN = (1.0, 0.85, 0.30)
SUN_EDGE = (0.90, 0.60, 0.10)
WALL = (0.98, 0.93, 0.80)
ROOF = (0.85, 0.42, 0.35)
DOOR = (0.55, 0.35, 0.20)
TRUNK = (0.55, 0.38, 0.22)
LEAF = (0.45, 0.72, 0.42)
PASTELS = (
    (0.87, 0.93, 1.00),
    (0.90, 0.97, 0.88),
    (1.00, 0.95, 0.84),
    (0.98, 0.89, 0.93),
    (0.93, 0.90, 0.98),
)


# --- spec --------------------------------------------------------------------------------------


def _check_text(value: str, field: str) -> str:
    """A single printable Latin-1 line that the pipeline will not mistake for an instruction."""
    value = value.strip()
    if not value:
        raise ValueError(f"{field}: text must not be empty")
    if "\n" in value or "\r" in value:
        raise ValueError(f"{field}: text must be a single line: {value!r}")
    try:
        value.encode("latin-1")
    except UnicodeEncodeError as error:
        raise ValueError(f"{field}: only Latin-1 characters render; got {value!r}") from error
    if INSTRUCTION_RE.match(value):
        raise ValueError(f"{field}: {value!r} would be read as an activity instruction")
    return value


class _Model(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)


class Picture(_Model):
    scene: str = Field(min_length=1, max_length=80)
    style: Literal["scene", "card"] = "scene"
    shapes: Annotated[list[ShapeName], Field(min_length=1, max_length=3)] | None = None

    @field_validator("scene")
    @classmethod
    def _scene(cls, value: str) -> str:
        return _check_text(value, "scene")

    @property
    def size(self) -> tuple[float, float]:
        return SCENE_SIZE if self.style == "scene" else CARD_SIZE

    @property
    def key(self) -> tuple[str, str, tuple[str, ...]]:
        return (self.style, self.scene, tuple(self.shapes or ()))


class Item(_Model):
    word: str = Field(min_length=1, max_length=24)
    picture: Picture | None = None

    @field_validator("word")
    @classmethod
    def _word(cls, value: str) -> str:
        return _check_text(value, "word")

    @property
    def card(self) -> Picture:
        return self.picture or Picture(scene=self.word, style="card")


class Activity(_Model):
    no: int = Field(ge=1, le=6)
    instruction: str = Field(min_length=1, max_length=60)
    type: ActivityType
    picture: Picture | None = None
    pattern: str | None = None
    items: list[Item] = Field(default_factory=list, max_length=2 * CARDS_PER_ROW)
    sound: str | None = None
    letters: str | None = None
    words: list[str] = Field(default_factory=list, max_length=8)
    options: list[str] = Field(default_factory=list, max_length=6)
    pairs: list[tuple[str, str]] = Field(default_factory=list, max_length=6)
    lines: list[str] = Field(default_factory=list, max_length=12)
    notes: str | None = None

    @field_validator("instruction")
    @classmethod
    def _instruction(cls, value: str) -> str:
        value = _check_text(value, "instruction")
        return value if value[-1] in ".!?" else value + "."

    @field_validator("pattern", "sound", "letters")
    @classmethod
    def _text(cls, value: str | None, info: ValidationInfo) -> str | None:
        return None if value is None else _check_text(value, str(info.field_name))

    @field_validator("words", "options", "lines")
    @classmethod
    def _texts(cls, values: list[str], info: ValidationInfo) -> list[str]:
        return [_check_text(value, str(info.field_name)) for value in values]

    @field_validator("pairs")
    @classmethod
    def _pairs(cls, values: list[tuple[str, str]]) -> list[tuple[str, str]]:
        return [(_check_text(left, "pairs"), _check_text(right, "pairs")) for left, right in values]


class Lesson(_Model):
    no: int = Field(ge=1, le=3)
    headings: Annotated[list[str], Field(min_length=2, max_length=2)] | None = None
    activities: list[Activity]
    notes: str | None = None

    @field_validator("headings")
    @classmethod
    def _headings(cls, values: list[str] | None) -> list[str] | None:
        return None if values is None else [_check_text(value, "headings") for value in values]

    @model_validator(mode="after")
    def _six_in_order(self) -> Lesson:
        numbers = [activity.no for activity in self.activities]
        if numbers != list(range(1, 7)):
            raise ValueError(
                f"lesson {self.no}: activities must be numbered 1 to 6 in order, got {numbers}"
            )
        return self


class Layout(_Model):
    lesson_line: bool = True


class UnitSpec(_Model):
    id: str = Field(pattern=r"^[a-z0-9][a-z0-9-]*$", max_length=40)
    grade: int = Field(ge=3, le=5)
    unit: int = Field(ge=1, le=20)
    title: str = Field(min_length=1, max_length=40)
    series: str = Field(default="classdeck-sample", min_length=1, max_length=40)
    layout: Layout = Field(default_factory=Layout)
    lessons: list[Lesson] = Field(min_length=1, max_length=3)
    notes: str | None = None

    @field_validator("title", "series")
    @classmethod
    def _title(cls, value: str, info: ValidationInfo) -> str:
        return _check_text(value, str(info.field_name))

    @model_validator(mode="after")
    def _lessons_in_order(self) -> UnitSpec:
        numbers = [lesson.no for lesson in self.lessons]
        if numbers != list(range(1, len(numbers) + 1)):
            raise ValueError(f"lessons must be numbered 1.. in order, got {numbers}")
        return self

    @property
    def heading(self) -> str:
        return f"Unit {self.unit}  {self.title}"

    @property
    def document_title(self) -> str:
        return f"Unit {self.unit} {self.title} ({self.series})"


def validate_spec(spec: dict | UnitSpec) -> UnitSpec:
    """The validated spec; raises ``pydantic.ValidationError`` (a ``ValueError``) otherwise."""
    return spec if isinstance(spec, UnitSpec) else UnitSpec.model_validate(spec)


def load_specs(units_dir: Path) -> list[dict]:
    """Every ``*.json`` under ``units_dir`` as a validated spec dict, in file-name order."""
    specs = []
    for path in sorted(Path(units_dir).glob("*.json")):
        try:
            spec = json.loads(path.read_text(encoding="utf-8"))
            unit = validate_spec(spec)
        except ValueError as error:
            raise ValueError(f"{path.name}: {error}") from error
        if unit.id != path.stem:
            raise ValueError(f"{path.name}: spec id {unit.id!r} must equal the file stem")
        specs.append(spec)
    return specs


# --- page plan (shared by the three renderers) -------------------------------------------------


@dataclass(frozen=True)
class Block:
    """One thing on a page: an instruction line, a text line, a picture or a row of cards."""

    kind: Literal["instruction", "text", "picture", "cards"]
    text: str = ""
    picture: Picture | None = None
    items: tuple[Item, ...] = ()


@dataclass(frozen=True)
class PagePlan:
    number: int
    lesson_no: int
    heading: str
    lesson_line: str | None
    blocks: tuple[Block, ...]


def plan_pages(spec: dict | UnitSpec) -> list[PagePlan]:
    """The pages of the unit in order: two per lesson, three activities each."""
    unit = validate_spec(spec)
    plans: list[PagePlan] = []
    for lesson in unit.lessons:
        headings = lesson.headings or [unit.heading, unit.heading]
        for half, heading in enumerate(headings):
            blocks: list[Block] = []
            for activity in lesson.activities[3 * half : 3 * half + 3]:
                activity_blocks = _activity_blocks(activity)
                for block in activity_blocks:
                    if block.kind != "cards" and _width(block.text) > CONTENT_WIDTH:
                        raise ValueError(
                            f"{unit.id}: lesson {lesson.no} activity {activity.no}: "
                            f"line too wide for the page: {block.text!r}"
                        )
                blocks.extend(activity_blocks)
            lesson_line = f"Lesson {lesson.no}" if unit.layout.lesson_line and half == 0 else None
            plans.append(PagePlan(len(plans) + 1, lesson.no, heading, lesson_line, tuple(blocks)))
    return plans


def _activity_blocks(activity: Activity) -> list[Block]:
    blocks = [Block("instruction", f"{activity.no}. {activity.instruction}")]
    if activity.pattern:
        blocks.append(Block("text", activity.pattern))
    if activity.picture:
        blocks.append(Block("picture", picture=activity.picture))
    for start in range(0, len(activity.items), CARDS_PER_ROW):
        row = tuple(activity.items[start : start + CARDS_PER_ROW])
        blocks.append(Block("cards", items=row))
        blocks.append(Block("text", _words_under_cards(row)))
    if activity.letters or activity.sound:
        blocks.append(Block("text", GAP.join(p for p in (activity.letters, activity.sound) if p)))
    blocks.extend(Block("text", text) for text in _fit_or_wrap(activity.words))
    blocks.extend(Block("text", text) for text in _fit_or_wrap(activity.options))
    blocks.extend(Block("text", text) for text in _pair_rows(activity.pairs))
    blocks.extend(Block("text", text) for text in activity.lines)
    return blocks


def _width(text: str, size: float = BODY_SIZE, font: str = BODY_FONT) -> float:
    return pymupdf.get_text_length(text, fontname=font, fontsize=size)


def _fit_or_wrap(texts: Sequence[str]) -> list[str]:
    """Short texts side by side on one line; one per line when they would not fit."""
    if not texts:
        return []
    joined = GAP.join(texts)
    return [joined] if _width(joined) <= CONTENT_WIDTH else list(texts)


def _pad_to(text: str, x: float) -> str:
    """``text`` followed by the spaces that bring the line to ``x`` points, at least one gap."""
    count = max(len(GAP), round((x - _width(text)) / _width(" ")))
    return text + " " * count


def _words_under_cards(items: Sequence[Item]) -> str:
    line = ""
    for index, item in enumerate(items):
        if index:
            line = _pad_to(line, index * (CARD_SIZE[0] + CARD_GAP))
        line += item.word
    return line


def _pair_rows(pairs: Sequence[tuple[str, str]]) -> list[str]:
    """Left column lettered a., b., ...; right column shifted one row so matching takes reading."""
    if not pairs:
        return []
    rights = [right for _, right in pairs]
    if len(rights) > 1:
        rights = rights[1:] + rights[:1]
    return [
        _pad_to(f"{chr(ord('a') + index)}. {left}", PAIR_COLUMN_X) + right
        for index, ((left, _), right) in enumerate(zip(pairs, rights, strict=True))
    ]


def _normalize(text: str) -> str:
    return " ".join(text.split())


def text_sequence(spec: dict | UnitSpec) -> list[list[tuple[str, str]]]:
    """Per page, the ``(kind, text)`` spans ingest is expected to produce from the PDF or DOCX.

    Text is whitespace-normalised the way ``classdeck.ingest`` does it. PPTX differs in one
    place: its ``Lesson L`` line is a ``paragraph``, since only the title placeholder is a
    heading there.
    """
    pages = []
    for plan in plan_pages(spec):
        spans = [("heading", _normalize(plan.heading))]
        if plan.lesson_line:
            spans.append(("heading", plan.lesson_line))
        for block in plan.blocks:
            if block.kind == "instruction":
                spans.append(("list_item", _normalize(block.text)))
            elif block.kind == "text":
                spans.append(("paragraph", _normalize(block.text)))
            elif block.kind == "picture":
                spans.append(("image", ""))
            else:
                spans.extend(("image", "") for _ in block.items)
        pages.append(spans)
    return pages


# --- pictures ----------------------------------------------------------------------------------

_SHAPE_KEYWORDS: tuple[tuple[str, frozenset[str]], ...] = (
    (
        "face",
        frozenset(
            [
                "child",
                "children",
                "girl",
                "boy",
                "friend",
                "friends",
                "pupil",
                "pupils",
                "teacher",
                "mother",
                "father",
                "family",
                "classmate",
                "classmates",
                "waving",
                "greeting",
                "people",
            ]
        ),
    ),
    (
        "sun",
        frozenset(
            [
                "sun",
                "sunny",
                "morning",
                "afternoon",
                "day",
                "outside",
                "park",
                "playground",
                "beach",
            ]
        ),
    ),
    (
        "house",
        frozenset(["house", "home", "school", "gate", "classroom", "door", "window", "room"]),
    ),
    ("ball", frozenset(["ball", "football", "play", "playing", "game", "toy", "toys"])),
    ("tree", frozenset(["tree", "trees", "garden", "forest"])),
)
_PLURAL_FACES = frozenset(["two", "children", "friends", "pupils", "classmates", "people"])


def infer_shapes(scene: str) -> list[str]:
    """Two or three shapes for a scene text, from its keywords; faces and a sun by default."""
    words = set(re.findall(r"[a-z]+", scene.lower()))
    shapes = [name for name, keys in _SHAPE_KEYWORDS if words & keys]
    if "face" in shapes and words & _PLURAL_FACES:
        shapes.insert(0, "face")
    for default in ("face", "sun", "house"):
        if len(shapes) >= 2:
            break
        if default not in shapes:
            shapes.append(default)
    return shapes[:3]


def _pastel(scene: str) -> tuple[float, float, float]:
    return PASTELS[zlib.crc32(scene.lower().encode("latin-1")) % len(PASTELS)]


def _face(shape: pymupdf.Shape, cx: float, cy: float, s: float) -> None:
    r = s * 0.42
    shape.draw_circle((cx, cy), r)
    shape.finish(color=INK, fill=SKIN, width=1.2)
    for side in (-1, 1):
        shape.draw_circle((cx + side * 0.38 * r, cy - 0.22 * r), 0.09 * r)
        shape.finish(color=INK, fill=INK, width=0.5)
    shape.draw_curve(
        (cx - 0.5 * r, cy + 0.3 * r), (cx, cy + 0.85 * r), (cx + 0.5 * r, cy + 0.3 * r)
    )
    shape.finish(color=INK, width=1.2, closePath=False)


def _sun(shape: pymupdf.Shape, cx: float, cy: float, s: float) -> None:
    r = s * 0.28
    for k in range(8):
        angle = math.tau * k / 8
        dx, dy = math.cos(angle), math.sin(angle)
        shape.draw_line(
            (cx + dx * r * 1.3, cy + dy * r * 1.3), (cx + dx * r * 1.75, cy + dy * r * 1.75)
        )
    shape.finish(color=SUN_EDGE, width=1.5, closePath=False)
    shape.draw_circle((cx, cy), r)
    shape.finish(color=SUN_EDGE, fill=SUN, width=1.2)


def _house(shape: pymupdf.Shape, cx: float, cy: float, s: float) -> None:
    half = s * 0.36
    shape.draw_rect(pymupdf.Rect(cx - half, cy - 0.05 * s, cx + half, cy + 0.5 * s))
    shape.finish(color=INK, fill=WALL, width=1.2)
    shape.draw_polyline(
        [(cx - half * 1.15, cy - 0.05 * s), (cx, cy - 0.5 * s), (cx + half * 1.15, cy - 0.05 * s)]
    )
    shape.finish(color=INK, fill=ROOF, width=1.2)
    shape.draw_rect(pymupdf.Rect(cx - 0.1 * s, cy + 0.2 * s, cx + 0.1 * s, cy + 0.5 * s))
    shape.finish(color=INK, fill=DOOR, width=1.0)


def _ball(shape: pymupdf.Shape, cx: float, cy: float, s: float) -> None:
    r = s * 0.36
    shape.draw_circle((cx, cy), r)
    shape.finish(color=INK, fill=WHITE, width=1.2)
    angles = (math.tau * k / 5 - math.pi / 2 for k in range(5))
    shape.draw_polyline(
        [(cx + math.cos(a) * r * 0.38, cy + math.sin(a) * r * 0.38) for a in angles]
    )
    shape.finish(color=INK, fill=INK, width=1.0)


def _tree(shape: pymupdf.Shape, cx: float, cy: float, s: float) -> None:
    shape.draw_rect(pymupdf.Rect(cx - 0.07 * s, cy + 0.05 * s, cx + 0.07 * s, cy + 0.5 * s))
    shape.finish(color=INK, fill=TRUNK, width=1.0)
    shape.draw_circle((cx, cy - 0.1 * s), s * 0.32)
    shape.finish(color=INK, fill=LEAF, width=1.2)


_DRAW: dict[str, Callable[[pymupdf.Shape, float, float, float], None]] = {
    "face": _face,
    "sun": _sun,
    "house": _house,
    "ball": _ball,
    "tree": _tree,
}


def _truncate(text: str, max_width: float, size: float) -> str:
    if _width(text, size) <= max_width:
        return text
    while text and _width(text + "...", size) > max_width:
        text = text[:-1]
    return text.rstrip() + "..."


def _draw_scene(page: pymupdf.Page, picture: Picture) -> None:
    width, height = picture.size
    shape = page.new_shape()
    shape.draw_rect(pymupdf.Rect(1, 1, width - 1, height - 1), radius=0.06)
    shape.finish(color=INK_SOFT, fill=_pastel(picture.scene), width=1.0)
    names = picture.shapes or infer_shapes(picture.scene)
    size = min(height * 0.5, width / (len(names) + 1) * 0.9)
    for index, name in enumerate(names):
        _DRAW[name](shape, width * (index + 1) / (len(names) + 1), height * 0.46, size)
    shape.commit()
    caption = _truncate(picture.scene, width - 16, CAPTION_SIZE)
    page.insert_text(
        (8, height - 7), caption, fontsize=CAPTION_SIZE, fontname=BODY_FONT, color=INK_SOFT
    )


def _card_shape(scene: str) -> str:
    """The one shape of a card: the first keyword match, else a stable pick from the text."""
    words = set(re.findall(r"[a-z]+", scene.lower()))
    for name, keys in _SHAPE_KEYWORDS:
        if words & keys:
            return name
    return SHAPES[zlib.crc32(scene.lower().encode("latin-1")) % len(SHAPES)]


def _draw_card(page: pymupdf.Page, picture: Picture) -> None:
    width, height = picture.size
    shape = page.new_shape()
    shape.draw_rect(pymupdf.Rect(1.5, 1.5, width - 1.5, height - 1.5), radius=0.1)
    shape.finish(color=INK_SOFT, fill=WHITE, width=1.2)
    names = picture.shapes or [_card_shape(picture.scene)]
    _DRAW[names[0]](shape, width / 2, height * 0.42, min(width, height) * 0.62)
    shape.commit()
    word = _truncate(picture.scene, width - 8, CARD_WORD_SIZE)
    x = (width - _width(word, CARD_WORD_SIZE)) / 2
    page.insert_text((x, height - 8), word, fontsize=CARD_WORD_SIZE, fontname=BODY_FONT, color=INK)


def draw_picture(picture: Picture) -> bytes:
    """The picture as PNG bytes: a vector drawing rendered to a pixmap, deterministic."""
    width, height = picture.size
    scratch = pymupdf.open()
    try:
        page = scratch.new_page(width=width, height=height)
        if picture.style == "scene":
            _draw_scene(page, picture)
        else:
            _draw_card(page, picture)
        return page.get_pixmap(dpi=PICTURE_DPI, alpha=False).tobytes("png")
    finally:
        scratch.close()


class _Pictures:
    """Draws each distinct picture once per unit and hands out its PNG bytes."""

    def __init__(self) -> None:
        self._cache: dict[tuple[str, str, tuple[str, ...]], bytes] = {}

    def png(self, picture: Picture) -> bytes:
        if picture.key not in self._cache:
            self._cache[picture.key] = draw_picture(picture)
        return self._cache[picture.key]


# --- PDF ---------------------------------------------------------------------------------------


def _render_pdf(plans: list[PagePlan], unit: UnitSpec, path: Path, pictures: _Pictures) -> None:
    document = pymupdf.open()
    try:
        document.set_metadata(
            {
                "title": unit.document_title,
                "author": AUTHOR,
                "subject": f"Grade {unit.grade} English, {unit.series}",
                "creationDate": PDF_DATE,
                "modDate": PDF_DATE,
            }
        )
        for plan in plans:
            page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
            _draw_page(page, plan, unit, pictures)
        document.save(str(path), garbage=4, deflate=True, no_new_id=True)
    finally:
        document.close()


def _draw_page(page: pymupdf.Page, plan: PagePlan, unit: UnitSpec, pictures: _Pictures) -> None:
    y = HEADING_BASELINE
    page.insert_text((MARGIN, y), plan.heading, fontsize=HEADING_SIZE, fontname=BOLD_FONT)
    y += HEADING_STEP
    if plan.lesson_line:
        page.insert_text((MARGIN, y), plan.lesson_line, fontsize=LESSON_SIZE, fontname=BOLD_FONT)
        y += LESSON_STEP
    for block in plan.blocks:
        if block.kind == "instruction":
            y += INSTRUCTION_GAP
            page.insert_text((MARGIN, y), block.text, fontsize=BODY_SIZE, fontname=BOLD_FONT)
            y += LINE_STEP
        elif block.kind == "text":
            page.insert_text((MARGIN, y), block.text, fontsize=BODY_SIZE, fontname=BODY_FONT)
            y += LINE_STEP
        elif block.kind == "picture" and block.picture is not None:
            width, height = block.picture.size
            top = y - PICTURE_RISE
            rect = pymupdf.Rect(MARGIN, top, MARGIN + width, top + height)
            page.insert_image(rect, stream=pictures.png(block.picture))
            y = top + height + PICTURE_DROP
        else:
            top = y - PICTURE_RISE
            for index, item in enumerate(block.items):
                x = MARGIN + index * (CARD_SIZE[0] + CARD_GAP)
                rect = pymupdf.Rect(x, top, x + CARD_SIZE[0], top + CARD_SIZE[1])
                page.insert_image(rect, stream=pictures.png(item.card))
            y = top + CARD_SIZE[1] + CARD_DROP
        if y > BOTTOM_LIMIT:
            raise ValueError(
                f"{unit.id}: page {plan.number} (lesson {plan.lesson_no}) overflows; "
                "use fewer lines or pictures"
            )


# --- DOCX --------------------------------------------------------------------------------------


def _render_docx(plans: list[PagePlan], unit: UnitSpec, path: Path, pictures: _Pictures) -> None:
    document = Document()
    _stamp(document.core_properties, unit)
    for index, plan in enumerate(plans):
        if index:
            document.add_page_break()
        document.add_heading(plan.heading, level=1)
        if plan.lesson_line:
            document.add_heading(plan.lesson_line, level=2)
        for block in plan.blocks:
            if block.kind == "instruction":
                _without_auto_number(document.add_paragraph(block.text, style="List Number"))
            elif block.kind == "text":
                document.add_paragraph(block.text)
            elif block.kind == "picture" and block.picture is not None:
                stream = io.BytesIO(pictures.png(block.picture))
                document.add_picture(stream, width=Inches(block.picture.size[0] / 72))
            else:
                paragraph = document.add_paragraph()
                for item in block.items:
                    stream = io.BytesIO(pictures.png(item.card))
                    paragraph.add_run().add_picture(stream, width=Inches(CARD_SIZE[0] / 72))
                    paragraph.add_run("  ")
    _save_office(document, path)


def _without_auto_number(paragraph: Paragraph) -> None:
    """Keep the "List Number" style but not its automatic counter: the text has its own number."""
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0


# --- PPTX --------------------------------------------------------------------------------------


def _render_pptx(plans: list[PagePlan], unit: UnitSpec, path: Path, pictures: _Pictures) -> None:
    presentation = Presentation()
    presentation.slide_width = PptxInches(SLIDE_WIDTH)
    presentation.slide_height = PptxInches(SLIDE_HEIGHT)
    _stamp(presentation.core_properties, unit)
    for plan in plans:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
        title = slide.shapes.title
        title.left, title.top = PptxInches(SLIDE_MARGIN), PptxInches(0.4)
        title.width, title.height = PptxInches(SLIDE_WIDTH - 2 * SLIDE_MARGIN), PptxInches(0.9)
        title.text = plan.heading
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(28)
        top = 1.5
        if plan.lesson_line:
            top = _add_text_box(slide, top, [(plan.lesson_line, True)], size=20)
        for kind, members in _segments(plan.blocks):
            if kind == "texts":
                lines = [(block.text, block.kind == "instruction") for block in members]
                top = _add_text_box(slide, top, lines, size=14)
            elif kind == "picture" and members[0].picture is not None:
                picture = members[0].picture
                stream = io.BytesIO(pictures.png(picture))
                width, height = picture.size
                slide.shapes.add_picture(
                    stream, PptxInches(SLIDE_MARGIN), PptxInches(top), width=PptxInches(width / 72)
                )
                top += height / 72 + 0.15
            else:
                for index, item in enumerate(members[0].items):
                    left = SLIDE_MARGIN + index * (CARD_SIZE[0] + CARD_GAP) / 72
                    stream = io.BytesIO(pictures.png(item.card))
                    slide.shapes.add_picture(
                        stream,
                        PptxInches(left),
                        PptxInches(top),
                        width=PptxInches(CARD_SIZE[0] / 72),
                    )
                top += CARD_SIZE[1] / 72 + 0.12
    _save_office(presentation, path)


def _segments(blocks: Sequence[Block]) -> list[tuple[str, list[Block]]]:
    """Consecutive text and instruction blocks become one text box; pictures stand alone."""
    segments: list[tuple[str, list[Block]]] = []
    for block in blocks:
        if block.kind in ("instruction", "text"):
            if segments and segments[-1][0] == "texts":
                segments[-1][1].append(block)
            else:
                segments.append(("texts", [block]))
        else:
            segments.append((block.kind, [block]))
    return segments


def _add_text_box(slide, top: float, lines: Sequence[tuple[str, bool]], *, size: int) -> float:
    height = SLIDE_LINE * len(lines)
    box = slide.shapes.add_textbox(
        PptxInches(SLIDE_MARGIN),
        PptxInches(top),
        PptxInches(SLIDE_WIDTH - 2 * SLIDE_MARGIN),
        PptxInches(height),
    )
    frame = box.text_frame
    frame.word_wrap = True
    for index, (text, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
    return top + height + 0.1


# --- shared output helpers ---------------------------------------------------------------------


def _stamp(properties, unit: UnitSpec) -> None:
    properties.title = unit.document_title
    properties.author = AUTHOR
    properties.last_modified_by = AUTHOR
    properties.subject = f"Grade {unit.grade} English, {unit.series}"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME


def _save_office(document, path: Path) -> None:
    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(_normalise_zip(buffer.getvalue()))


def _normalise_zip(data: bytes) -> bytes:
    """The archive with every entry stamped 1980-01-01, so two saves give identical bytes."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            stamped = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE_TIME)
            stamped.compress_type = info.compress_type
            stamped.external_attr = info.external_attr
            target.writestr(stamped, source.read(info))
    return out.getvalue()


_RENDERERS: dict[str, Callable[[list[PagePlan], UnitSpec, Path, _Pictures], None]] = {
    "pdf": _render_pdf,
    "docx": _render_docx,
    "pptx": _render_pptx,
}


def render_unit(
    spec: dict | UnitSpec, out_dir: Path, formats: Sequence[str] = FORMATS
) -> dict[str, Path]:
    """Render ``spec`` to ``out_dir / <id>.<format>`` for each format; returns the paths."""
    unit = validate_spec(spec)
    unknown = [fmt for fmt in formats if fmt not in FORMATS]
    if unknown:
        raise ValueError(f"unknown formats {unknown}; choose from {list(FORMATS)}")
    plans = plan_pages(unit)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pictures = _Pictures()
    written: dict[str, Path] = {}
    for fmt in formats:
        path = out_dir / f"{unit.id}.{fmt}"
        _RENDERERS[fmt](plans, unit, path, pictures)
        written[fmt] = path
    return written
